"""Stage 13: the norms/tariffs/rule-parameters exports are the lists on
paper — same filters, readable cells, the id last. All three registers are
open to any authenticated user (no zone of their own), so the "no scope"
shape here is simply: a different, unrelated authenticated caller sees the
SAME rows the export shows to anyone else."""

import io
import uuid

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.modules.gis.models import Contour

pytestmark = pytest.mark.asyncio


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def _draft_norm(
    client: AsyncClient, contour_id: uuid.UUID, activity_id: uuid.UUID, **over
) -> dict:
    payload = {
        "contour_id": str(contour_id),
        "activity_type_id": str(activity_id),
        "yield_c_per_ha": "12.0",
        "season": {"windows": [{"from": "04-01", "to": "10-31"}]},
        "rotation": {"rest_years": []},
        "effective_from": "2030-01-01",
    } | over
    created = await client.post("/api/v1/norms", json=payload)
    assert created.status_code == 201, created.text
    return created.json()


# --- norms ---------------------------------------------------------------


async def test_norms_export_holds_exactly_the_rows_the_list_shows(
    gis_specialist_client: AsyncClient,
    published_contour: Contour,
    grazing_activity_id: uuid.UUID,
) -> None:
    await _draft_norm(gis_specialist_client, published_contour.id, grazing_activity_id)

    listed = await gis_specialist_client.get(f"/api/v1/norms?contour_id={published_contour.id}")
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await gis_specialist_client.get(
        "/api/v1/norms/export.xlsx",
        params={"contour_id": str(published_contour.id), "lang": "ru"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Контур" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_norms_export_applies_the_same_filter_as_the_list(
    gis_specialist_client: AsyncClient,
    published_contour: Contour,
    grazing_activity_id: uuid.UUID,
    haymaking_activity_id: uuid.UUID,
) -> None:
    await _draft_norm(gis_specialist_client, published_contour.id, grazing_activity_id)

    resp = await gis_specialist_client.get(
        "/api/v1/norms/export.xlsx",
        params={
            "contour_id": str(published_contour.id),
            "activity_type_id": str(haymaking_activity_id),
        },
    )
    assert resp.status_code == 200, resp.text
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_norms_export_renders_labels_not_codes(
    gis_specialist_client: AsyncClient,
    published_contour: Contour,
    grazing_activity_id: uuid.UUID,
) -> None:
    await _draft_norm(gis_specialist_client, published_contour.id, grazing_activity_id)

    resp = await gis_specialist_client.get(
        "/api/v1/norms/export.xlsx",
        params={"contour_id": str(published_contour.id), "lang": "uz_latn"},
    )
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == published_contour.number  # the contour number first
    assert row[2] == "Qoralama"  # the status label, not "draft"


async def test_norms_export_truncates_at_the_cap_and_says_so(
    gis_specialist_client: AsyncClient,
    published_contour: Contour,
    grazing_activity_id: uuid.UUID,
    haymaking_activity_id: uuid.UUID,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await _draft_norm(gis_specialist_client, published_contour.id, grazing_activity_id)
    await _draft_norm(gis_specialist_client, published_contour.id, haymaking_activity_id)

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db_arg, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_arg, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await gis_specialist_client.get(
        "/api/v1/norms/export.xlsx", params={"contour_id": str(published_contour.id)}
    )
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_norms_export_is_visible_to_any_authenticated_caller(
    gis_specialist_client: AsyncClient,
    applicant_client: AsyncClient,
    published_contour: Contour,
    grazing_activity_id: uuid.UUID,
) -> None:
    """No zone of its own (`router.py`'s module docstring): a caller with no
    special grant sees the SAME rows the list shows it — the export must
    not narrow (or widen) that."""
    await _draft_norm(gis_specialist_client, published_contour.id, grazing_activity_id)

    listed = await applicant_client.get(f"/api/v1/norms?contour_id={published_contour.id}")
    assert listed.status_code == 200
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await applicant_client.get(
        "/api/v1/norms/export.xlsx", params={"contour_id": str(published_contour.id)}
    )
    assert resp.status_code == 200, resp.text
    exported_ids = {
        str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
    }
    assert exported_ids == listed_ids


async def test_norms_export_rejects_an_unknown_language(
    gis_specialist_client: AsyncClient,
) -> None:
    resp = await gis_specialist_client.get("/api/v1/norms/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- tariffs ---------------------------------------------------------------
# The seeded VMQ 278 rates (`test_tariffs_api.py`'s own fixtures) are already
# in force from 2015-09-30 — haymaking (1.5) and grazing's four livestock
# groups — so these tests read them rather than creating new ones.


async def test_tariffs_export_holds_exactly_the_rows_the_list_shows(
    applicant_client: AsyncClient,
) -> None:
    listed = await applicant_client.get(
        "/api/v1/tariffs", params={"activity_code": "grazing", "on_date": "2026-08-30"}
    )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await applicant_client.get(
        "/api/v1/tariffs/export.xlsx",
        params={"activity_code": "grazing", "on_date": "2026-08-30", "lang": "ru"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Вид деятельности" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_tariffs_export_applies_the_same_filter_as_the_list(
    applicant_client: AsyncClient,
) -> None:
    """A date before any rate took effect: empty on both the list and the
    export (`test_a_tariff_for_a_date_before_it_took_effect_is_not_returned`'s
    own scenario)."""
    resp = await applicant_client.get(
        "/api/v1/tariffs/export.xlsx",
        params={"activity_code": "haymaking", "on_date": "2015-01-01"},
    )
    assert resp.status_code == 200, resp.text
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_tariffs_export_renders_labels_not_codes(applicant_client: AsyncClient) -> None:
    resp = await applicant_client.get(
        "/api/v1/tariffs/export.xlsx",
        params={"activity_code": "grazing", "on_date": "2026-08-30", "lang": "uz_latn"},
    )
    rows = {row[1]: row for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)}
    assert "Yirik chorva, katta" in rows  # the large_adult label, not the code
    row = rows["Yirik chorva, katta"]
    assert row[5] == "Eʼlon qilingan"  # the status label, not "published"
    assert row[3] == "bosh"  # the quantity_unit label, not "head"


async def test_tariffs_export_truncates_at_the_cap_and_says_so(
    applicant_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db_arg, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_arg, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await applicant_client.get(
        "/api/v1/tariffs/export.xlsx",
        params={"activity_code": "grazing", "on_date": "2026-08-30"},
    )
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2  # grazing has four published groups
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_tariffs_export_is_visible_to_any_authenticated_caller(
    applicant_client: AsyncClient,
) -> None:
    """No zone of its own: a plain applicant sees the same in-force rates
    the list shows anyone."""
    listed = await applicant_client.get(
        "/api/v1/tariffs", params={"activity_code": "haymaking", "on_date": "2026-08-30"}
    )
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await applicant_client.get(
        "/api/v1/tariffs/export.xlsx",
        params={"activity_code": "haymaking", "on_date": "2026-08-30"},
    )
    assert resp.status_code == 200, resp.text
    exported_ids = {
        str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
    }
    assert exported_ids == listed_ids


async def test_tariffs_export_rejects_an_unknown_language(applicant_client: AsyncClient) -> None:
    resp = await applicant_client.get("/api/v1/tariffs/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- rule parameters ---------------------------------------------------------


async def _draft_parameter(client: AsyncClient, code: str, value: str = "0.9") -> dict:
    created = await client.post(
        "/api/v1/rule-parameters",
        json={
            "code": code,
            "value": value,
            "effective_from": "2030-01-01",
            "basis": "test",
        },
    )
    assert created.status_code == 201, created.text
    return created.json()


async def test_parameters_export_holds_exactly_the_rows_the_list_shows(
    tariffs_maker_client: AsyncClient, unique_suffix: str
) -> None:
    code = f"test_param_{unique_suffix}"
    await _draft_parameter(tariffs_maker_client, code)

    listed = await tariffs_maker_client.get("/api/v1/rule-parameters", params={"code": code})
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await tariffs_maker_client.get(
        "/api/v1/rule-parameters/export.xlsx", params={"code": code, "lang": "ru"}
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Код" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_parameters_export_applies_the_same_filter_as_the_list(
    tariffs_maker_client: AsyncClient, unique_suffix: str
) -> None:
    code = f"test_param_{unique_suffix}"
    await _draft_parameter(tariffs_maker_client, code)

    resp = await tariffs_maker_client.get(
        "/api/v1/rule-parameters/export.xlsx", params={"code": f"other_{unique_suffix}"}
    )
    assert resp.status_code == 200, resp.text
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_parameters_export_renders_labels_not_codes(
    tariffs_maker_client: AsyncClient, unique_suffix: str
) -> None:
    code = f"test_param_{unique_suffix}"
    await _draft_parameter(tariffs_maker_client, code)

    resp = await tariffs_maker_client.get(
        "/api/v1/rule-parameters/export.xlsx", params={"code": code, "lang": "uz_latn"}
    )
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == code  # the code first
    assert row[3] == "Qoralama"  # the status label, not "draft"
    assert row[7] == "Test User"  # created_by resolved to a name (make_user's default), not a uuid


async def test_parameters_export_truncates_at_the_cap_and_says_so(
    tariffs_maker_client: AsyncClient, unique_suffix: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    await _draft_parameter(tariffs_maker_client, f"test_param_a_{unique_suffix}")
    await _draft_parameter(tariffs_maker_client, f"test_param_b_{unique_suffix}")

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db_arg, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_arg, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    # No `code` filter: the shared, persistent test DB always carries more
    # than one parameter (the seeded `coef_sb:*`/`bhm`/... rows, plus the two
    # just created), so `total` is guaranteed > the cap of 1 here.
    resp = await tariffs_maker_client.get("/api/v1/rule-parameters/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total > 1
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_parameters_export_is_visible_to_any_authenticated_caller(
    tariffs_maker_client: AsyncClient, applicant_client: AsyncClient, unique_suffix: str
) -> None:
    code = f"test_param_{unique_suffix}"
    await _draft_parameter(tariffs_maker_client, code)

    listed = await applicant_client.get("/api/v1/rule-parameters", params={"code": code})
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await applicant_client.get("/api/v1/rule-parameters/export.xlsx", params={"code": code})
    assert resp.status_code == 200, resp.text
    exported_ids = {
        str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
    }
    assert exported_ids == listed_ids


async def test_parameters_export_rejects_an_unknown_language(
    tariffs_maker_client: AsyncClient,
) -> None:
    resp = await tariffs_maker_client.get(
        "/api/v1/rule-parameters/export.xlsx", params={"lang": "en"}
    )
    assert resp.status_code == 422
