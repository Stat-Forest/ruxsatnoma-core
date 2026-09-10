"""Stage 13 (ruling #204): the admin module's registers on paper — users,
organizations, announcements, legal documents, and the integrations
outbox/DLQ. Same scope, same filters, readable cells, the id last."""

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.core import settings_store
from app.main import create_app
from app.modules.admin.models import Organization
from app.modules.admin.permissions import (
    ANNOUNCEMENTS_MANAGE,
    INTEGRATIONS_VIEW,
    LEGAL_DOCUMENTS_MANAGE,
)
from app.modules.auth.permissions import USERS_MANAGE, USERS_VIEW
from app.modules.integrations import senders
from app.modules.integrations import service as integrations_service
from tests.conftest import make_client
from tests.modules.admin.test_announcements import make_announcement
from tests.modules.admin.test_organizations_admin import auth_client, signed_in_with
from tests.modules.auth.test_sessions import make_user
from tests.modules.gis.conftest import leshoz as leshoz  # noqa: F401
from tests.modules.gis.conftest import other_leshoz as other_leshoz  # noqa: F401

pytestmark = pytest.mark.asyncio

API = "/api/v1"
LEGAL_DOCUMENT_BODY = {
    "title": {"uz_latn": "Test hujjat", "ru": "Тестовый акт"},
    "doc_number": "TEST-1",
    "adopted_on": "2020-01-01",
    "source_url": "https://lex.uz/docs/test",
}


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None  # a workbook this module wrote always has one
    return sheet


def _exported_ids(content: bytes) -> set[str]:
    return {str(row[-1]) for row in _sheet(content).iter_rows(min_row=2, values_only=True)}


# ---------------------------------------------------------------------------
# Users — GET /admin/users/export.xlsx
# ---------------------------------------------------------------------------


async def test_users_export_holds_exactly_the_rows_the_list_shows(db, leshoz):  # noqa: F811
    manager, token, csrf = await signed_in_with(db, USERS_MANAGE)
    target = await make_user(db, organization_id=leshoz.id)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/users", params={"q": target.login})
        resp = await client.get(
            f"{API}/admin/users/export.xlsx", params={"q": target.login, "lang": "ru"}
        )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids
    assert exported_ids == {str(target.id)}


async def test_users_export_applies_the_same_filters_as_the_list(db, leshoz, other_leshoz):  # noqa: F811
    manager, token, csrf = await signed_in_with(db, USERS_MANAGE)
    in_org = await make_user(db, organization_id=leshoz.id)
    await make_user(db, organization_id=other_leshoz.id)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/users/export.xlsx", params={"organization_id": str(leshoz.id)}
        )
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert str(in_org.id) in exported_ids


async def test_users_export_renders_labels_not_codes(db, leshoz):  # noqa: F811
    manager, token, csrf = await signed_in_with(db, USERS_MANAGE)
    target = await make_user(db, organization_id=leshoz.id)
    target.full_name = "Aaaaa Export Target"
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/users/export.xlsx", params={"q": target.login, "lang": "uz_latn"}
        )
    assert resp.status_code == 200, resp.text
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == "Aaaaa Export Target"  # F.I.Sh. first
    assert row[6] == "Faol"  # status label, not "active"


async def test_users_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    manager, token, csrf = await signed_in_with(db, USERS_MANAGE)
    await make_user(db)
    await make_user(db)
    await db.commit()

    real_get_int = settings_store.get_int

    async def one(_db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(_db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/users/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_users_export_without_the_permission_matches_the_list_status(db):
    _, token, csrf = await signed_in_with(db)  # no grants
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/users")
        resp = await client.get(f"{API}/admin/users/export.xlsx")
    assert listed.status_code == 403
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_users_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db, USERS_VIEW)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/users/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Organizations — GET /refs/organizations/export.xlsx
# ---------------------------------------------------------------------------


async def test_organizations_export_holds_exactly_the_rows_the_list_shows(db, agency):
    _, token, csrf = await signed_in_with(db)  # authenticated, no permission needed (ruling 10)
    suffix = uuid.uuid4().hex[:8]
    parent = Organization(
        kind="territorial", code=f"parent-{suffix}", name={"uz_cyrl": "Тест"}, parent_id=agency.id
    )
    db.add(parent)
    await db.flush()
    child = Organization(
        kind="leshoz", code=f"child-{suffix}", name={"uz_cyrl": "Бола"}, parent_id=parent.id
    )
    db.add(child)
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/refs/organizations", params={"parent_id": str(parent.id)})
        resp = await client.get(
            f"{API}/refs/organizations/export.xlsx",
            params={"parent_id": str(parent.id), "lang": "ru"},
        )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids == {str(child.id)}


async def test_organizations_export_applies_the_same_filters_as_the_list(db, agency):
    _, token, csrf = await signed_in_with(db)
    suffix = uuid.uuid4().hex[:8]
    parent = Organization(
        kind="territorial", code=f"parent2-{suffix}", name={"uz_cyrl": "Тест2"}, parent_id=agency.id
    )
    db.add(parent)
    await db.flush()
    active_child = Organization(
        kind="leshoz",
        code=f"active-{suffix}",
        name={"uz_cyrl": "Актив"},
        parent_id=parent.id,
        status="active",
    )
    archived_child = Organization(
        kind="leshoz",
        code=f"archived-{suffix}",
        name={"uz_cyrl": "Архив"},
        parent_id=parent.id,
        status="archived",
    )
    db.add_all([active_child, archived_child])
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/refs/organizations/export.xlsx",
            params={"parent_id": str(parent.id), "status": "archived"},
        )
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == {str(archived_child.id)}


async def test_organizations_export_renders_labels_not_codes(db, agency):
    _, token, csrf = await signed_in_with(db)
    suffix = uuid.uuid4().hex[:8]
    parent = Organization(
        kind="territorial", code=f"parent3-{suffix}", name={"uz_cyrl": "Тест3"}, parent_id=agency.id
    )
    db.add(parent)
    await db.flush()
    child = Organization(
        kind="leshoz",
        code=f"leshoz-{suffix}",
        name={"uz_latn": "A Test Leshoz"},
        parent_id=parent.id,
    )
    db.add(child)
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/refs/organizations/export.xlsx",
            params={"parent_id": str(parent.id), "lang": "uz_latn"},
        )
    assert resp.status_code == 200, resp.text
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == "A Test Leshoz"  # name first
    assert row[2] == "Oʻrmon xoʻjaligi"  # kind label, not "leshoz"


async def test_organizations_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await signed_in_with(db)
    await db.commit()

    real_get_int = settings_store.get_int

    async def one(_db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(_db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/refs/organizations/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_organizations_export_without_authentication_matches_the_list_status(db):
    async with make_client(create_app(), lifespan=True) as client:
        listed = await client.get(f"{API}/refs/organizations")
        resp = await client.get(f"{API}/refs/organizations/export.xlsx")
    assert listed.status_code == 401
    assert resp.status_code == 401
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_organizations_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/refs/organizations/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Announcements — GET /admin/announcements/export.xlsx
# ---------------------------------------------------------------------------


async def test_announcements_export_holds_exactly_the_rows_the_list_shows(db):
    manager, token, csrf = await signed_in_with(db, ANNOUNCEMENTS_MANAGE)
    target = await make_announcement(
        db,
        created_by=manager.id,
        status="draft",
        title={"uz_cyrl": f"Экспорт {uuid.uuid4().hex[:8]}"},
    )
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(
            f"{API}/admin/announcements", params={"status": "draft", "page_size": 100}
        )
        resp = await client.get(
            f"{API}/admin/announcements/export.xlsx", params={"status": "draft", "lang": "ru"}
        )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids
    assert str(target.id) in exported_ids


async def test_announcements_export_applies_the_same_filters_as_the_list(db):
    manager, token, csrf = await signed_in_with(db, ANNOUNCEMENTS_MANAGE)
    marker = uuid.uuid4().hex[:8]
    draft = await make_announcement(
        db, created_by=manager.id, status="draft", title={"uz_cyrl": f"Д{marker}"}
    )
    published = await make_announcement(
        db, created_by=manager.id, status="published", title={"uz_cyrl": f"П{marker}"}
    )
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/announcements/export.xlsx", params={"status": "published"}
        )
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert str(published.id) in exported_ids
    assert str(draft.id) not in exported_ids


async def test_announcements_export_renders_labels_not_codes(db):
    manager, token, csrf = await signed_in_with(db, ANNOUNCEMENTS_MANAGE)
    marker = uuid.uuid4().hex[:8]
    await make_announcement(
        db,
        created_by=manager.id,
        status="draft",
        title={"uz_latn": f"Export title {marker}"},
        audience=None,
    )
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/announcements/export.xlsx", params={"status": "draft", "lang": "uz_latn"}
        )
    assert resp.status_code == 200, resp.text
    rows = list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    row = next(r for r in rows if r[0] == f"Export title {marker}")
    assert row[1] == "Qoralama"  # status label, not "draft"
    assert row[2] == "Barcha foydalanuvchilar"  # no audience -> everyone


async def test_announcements_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    manager, token, csrf = await signed_in_with(db, ANNOUNCEMENTS_MANAGE)
    await make_announcement(db, created_by=manager.id, status="draft")
    await make_announcement(db, created_by=manager.id, status="draft")
    await db.commit()

    real_get_int = settings_store.get_int

    async def one(_db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(_db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/announcements/export.xlsx", params={"status": "draft"}
        )
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_announcements_export_without_the_permission_matches_the_list_status(db):
    _, token, csrf = await signed_in_with(db)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/announcements")
        resp = await client.get(f"{API}/admin/announcements/export.xlsx")
    assert listed.status_code == 403
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_announcements_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db, ANNOUNCEMENTS_MANAGE)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/announcements/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Legal documents — GET /admin/legal-documents/export.xlsx
# ---------------------------------------------------------------------------


async def test_legal_documents_export_holds_exactly_the_rows_the_list_shows(db):
    _, token, csrf = await signed_in_with(db, LEGAL_DOCUMENTS_MANAGE)
    await db.commit()
    marker = uuid.uuid4().hex[:8]
    body = {**LEGAL_DOCUMENT_BODY, "doc_number": f"TEST-{marker}"}
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(f"{API}/admin/legal-documents", json=body)
        assert created.status_code == 201, created.text
        doc_id = created.json()["id"]
        listed = await client.get(
            f"{API}/admin/legal-documents", params={"status": "draft", "page_size": 100}
        )
        resp = await client.get(
            f"{API}/admin/legal-documents/export.xlsx", params={"status": "draft", "lang": "ru"}
        )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids
    assert doc_id in exported_ids


async def test_legal_documents_export_applies_the_same_filters_as_the_list(db):
    _, token, csrf = await signed_in_with(db, LEGAL_DOCUMENTS_MANAGE)
    await db.commit()
    marker = uuid.uuid4().hex[:8]
    body = {**LEGAL_DOCUMENT_BODY, "doc_number": f"PUB-{marker}"}
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(f"{API}/admin/legal-documents", json=body)
        doc_id = created.json()["id"]
        published = await client.post(f"{API}/admin/legal-documents/{doc_id}/publish")
        assert published.status_code == 200, published.text
        resp = await client.get(
            f"{API}/admin/legal-documents/export.xlsx", params={"status": "published"}
        )
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert doc_id in exported_ids


async def test_legal_documents_export_renders_labels_not_codes(db):
    _, token, csrf = await signed_in_with(db, LEGAL_DOCUMENTS_MANAGE)
    await db.commit()
    marker = uuid.uuid4().hex[:8]
    body = {**LEGAL_DOCUMENT_BODY, "doc_number": f"LBL-{marker}"}
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(f"{API}/admin/legal-documents", json=body)
        doc_id = created.json()["id"]
        resp = await client.get(
            f"{API}/admin/legal-documents/export.xlsx",
            params={"status": "draft", "lang": "uz_latn"},
        )
    assert resp.status_code == 200, resp.text
    rows = list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    row = next(r for r in rows if r[-1] == doc_id)
    assert row[0] == f"LBL-{marker}"  # doc_number first
    assert row[4] == "Qoralama"  # status label
    assert row[3] == "lex.uz havolasi"  # source label from source_url


async def test_legal_documents_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await signed_in_with(db, LEGAL_DOCUMENTS_MANAGE)
    await db.commit()
    marker = uuid.uuid4().hex[:8]
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        await client.post(
            f"{API}/admin/legal-documents",
            json={**LEGAL_DOCUMENT_BODY, "doc_number": f"CAP1-{marker}"},
        )
        await client.post(
            f"{API}/admin/legal-documents",
            json={**LEGAL_DOCUMENT_BODY, "doc_number": f"CAP2-{marker}"},
        )

        real_get_int = settings_store.get_int

        async def one(_db, key):
            if key == "register_export_max_rows":
                return 1
            return await real_get_int(_db, key)

        monkeypatch.setattr(settings_store, "get_int", one)
        resp = await client.get(
            f"{API}/admin/legal-documents/export.xlsx", params={"status": "draft"}
        )
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_legal_documents_export_without_the_permission_matches_the_list_status(db):
    _, token, csrf = await signed_in_with(db)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/legal-documents")
        resp = await client.get(f"{API}/admin/legal-documents/export.xlsx")
    assert listed.status_code == 403
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_legal_documents_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db, LEGAL_DOCUMENTS_MANAGE)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/legal-documents/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Integrations outbox — GET /admin/integrations/outbox/export.xlsx
# ---------------------------------------------------------------------------


async def test_outbox_export_holds_exactly_the_rows_the_list_shows(db):
    _, token, csrf = await signed_in_with(db, INTEGRATIONS_VIEW)
    await db.commit()
    destination = f"_export_test_{uuid.uuid4().hex[:8]}"
    msg = await integrations_service.enqueue(db, destination=destination, payload={"x": 1})
    assert msg is not None
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(
            f"{API}/admin/integrations/outbox", params={"destination": destination}
        )
        resp = await client.get(
            f"{API}/admin/integrations/outbox/export.xlsx",
            params={"destination": destination, "lang": "ru"},
        )
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids == {str(msg.id)}


async def test_outbox_export_applies_the_same_filters_as_the_list(db):
    _, token, csrf = await signed_in_with(db, INTEGRATIONS_VIEW)
    await db.commit()
    ok_destination = f"_export_ok_{uuid.uuid4().hex[:8]}"

    async def ok_sender(_db, _payload: dict) -> None:
        return None

    senders.SENDERS[ok_destination] = ok_sender
    try:
        delivered = await integrations_service.enqueue(db, destination=ok_destination, payload={})
        assert delivered is not None
        await db.commit()
        assert await integrations_service.deliver_one(db) is True

        pending_destination = f"_export_pending_{uuid.uuid4().hex[:8]}"
        pending = await integrations_service.enqueue(
            db, destination=pending_destination, payload={}
        )
        assert pending is not None
        await db.commit()
    finally:
        senders.SENDERS.pop(ok_destination, None)

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/integrations/outbox/export.xlsx",
            params={"status": "pending", "destination": pending_destination},
        )
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == {str(pending.id)}


async def test_outbox_export_renders_labels_not_codes_and_never_the_payload(db):
    _, token, csrf = await signed_in_with(db, INTEGRATIONS_VIEW)
    await db.commit()
    destination = f"_export_label_{uuid.uuid4().hex[:8]}"
    msg = await integrations_service.enqueue(
        db, destination=destination, payload={"secret": "should-never-appear"}
    )
    assert msg is not None
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/integrations/outbox/export.xlsx",
            params={"destination": destination, "lang": "uz_latn"},
        )
    assert resp.status_code == 200, resp.text
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == destination
    assert row[1] == "Navbatda"  # status label, not "pending"
    assert "should-never-appear" not in "".join(str(v) for v in row if v is not None)


async def test_outbox_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await signed_in_with(db, INTEGRATIONS_VIEW)
    await db.commit()
    d1 = f"_export_cap1_{uuid.uuid4().hex[:8]}"
    d2 = f"_export_cap2_{uuid.uuid4().hex[:8]}"
    await integrations_service.enqueue(db, destination=d1, payload={})
    await integrations_service.enqueue(db, destination=d2, payload={})
    await db.commit()

    real_get_int = settings_store.get_int

    async def one(_db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(_db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/integrations/outbox/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_outbox_export_without_the_permission_matches_the_list_status(db):
    _, token, csrf = await signed_in_with(db)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/integrations/outbox")
        resp = await client.get(f"{API}/admin/integrations/outbox/export.xlsx")
    assert listed.status_code == 403
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_outbox_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db, INTEGRATIONS_VIEW)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/integrations/outbox/export.xlsx", params={"lang": "en"}
        )
    assert resp.status_code == 422
