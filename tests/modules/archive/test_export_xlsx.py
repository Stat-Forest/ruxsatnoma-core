"""Stage 13: `GET /archive/export.xlsx` — the archive register on paper.
Same zone, same filters, readable cells, the id last."""

import io

from openpyxl import load_workbook

from app.modules.archive.permissions import ARCHIVE_MANAGE, ARCHIVE_VIEW
from tests.modules.archive.conftest import _client_for, _client_with_role, make_application


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def test_export_holds_exactly_the_rows_the_list_shows(db, leshoz, other_leshoz):
    mine = await make_application(db, org=leshoz, status="CLOSED", applicant_name="A A")
    theirs = await make_application(db, org=other_leshoz, status="CLOSED", applicant_name="B B")
    await db.commit()

    async for client in _client_for(db, ARCHIVE_MANAGE, organization_id=leshoz.id):
        archived_mine = await client.post(f"/api/v1/archive/application/{mine.id}", json={})
        assert archived_mine.status_code == 200

    async for client in _client_for(db, ARCHIVE_MANAGE, organization_id=other_leshoz.id):
        archived_theirs = await client.post(f"/api/v1/archive/application/{theirs.id}", json={})
        assert archived_theirs.status_code == 200

    async for client in _client_for(db, ARCHIVE_VIEW, organization_id=leshoz.id):
        listed = await client.get("/api/v1/archive", params={"page_size": 100})
        listed_ids = {row["object_id"] for row in listed.json()["items"]}
        assert str(theirs.id) not in listed_ids  # the zone hides it from the screen

        resp = await client.get("/api/v1/archive/export.xlsx", params={"lang": "ru"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert resp.headers["x-export-truncated"] == "false"
        sheet = _sheet(resp.content)
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "Тип объекта" and headers[-1] == "ID"
        exported_object_ids = {str(row[1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert str(mine.id) in exported_object_ids
        assert exported_object_ids == listed_ids


async def test_export_applies_the_same_filters_as_the_list(db, leshoz):
    application = await make_application(db, org=leshoz, status="CANCELLED", applicant_name="A A")
    await db.commit()

    async for client in _client_for(db, ARCHIVE_MANAGE, organization_id=leshoz.id):
        archived = await client.post(f"/api/v1/archive/application/{application.id}", json={})
        assert archived.status_code == 200

    async for client in _client_for(db, ARCHIVE_VIEW, organization_id=leshoz.id):
        resp = await client.get(
            "/api/v1/archive/export.xlsx", params={"object_type": "permit", "lang": "uz_latn"}
        )
        assert resp.status_code == 200
        assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_export_renders_labels_not_codes(db, leshoz):
    application = await make_application(db, org=leshoz, status="CLOSED", applicant_name="A A")
    await db.commit()

    async for client in _client_for(db, ARCHIVE_MANAGE, organization_id=leshoz.id):
        archived = await client.post(f"/api/v1/archive/application/{application.id}", json={})
        assert archived.status_code == 200

    async for client in _client_for(db, ARCHIVE_VIEW, organization_id=leshoz.id):
        resp = await client.get("/api/v1/archive/export.xlsx", params={"lang": "uz_latn"})
        row = next(
            r
            for r in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
            if r[1] == str(application.id)
        )
        assert row[0] == "Ariza"  # the object_type LABEL, not "application"
        assert row[3] == "Saqlangan"  # the status LABEL, not "stored"


async def test_export_truncates_at_the_cap_and_says_so(db, leshoz, monkeypatch):
    for i in range(3):
        app = await make_application(db, org=leshoz, status="CLOSED", applicant_name=f"Cap {i}")
        await db.commit()
        async for client in _client_for(db, ARCHIVE_MANAGE, organization_id=leshoz.id):
            resp = await client.post(f"/api/v1/archive/application/{app.id}", json={})
            assert resp.status_code == 200

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def capped(db_, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_, key)

    monkeypatch.setattr(settings_store, "get_int", capped)

    async for client in _client_for(db, ARCHIVE_VIEW, organization_id=leshoz.id):
        resp = await client.get("/api/v1/archive/export.xlsx")
        assert resp.status_code == 200
        total = int(resp.headers["x-export-total"])
        assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
        assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1
        assert resp.headers["x-export-rows"] == str(min(total, 1))


async def test_export_requires_the_permission(db, leshoz):
    async for client in _client_with_role(db, "inspector"):
        resp = await client.get("/api/v1/archive/export.xlsx")
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_export_rejects_an_unknown_language(db, leshoz):
    async for client in _client_for(db, ARCHIVE_VIEW, organization_id=leshoz.id):
        resp = await client.get("/api/v1/archive/export.xlsx", params={"lang": "en"})
        assert resp.status_code == 422
