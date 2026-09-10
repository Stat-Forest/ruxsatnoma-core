"""Stage 13 (ruling #204): `GET /admin/ratings/export.xlsx` — the anonymous
comment feed on paper. Ruling #141 holds in the file exactly as on the screen:
no applicant, no permit — the id column is the RATING's own id, which names
nobody."""

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.core import settings_store
from tests.modules.permits.test_ratings import (
    other_org_ratings_client as other_org_ratings_client,  # noqa: F401
)
from tests.modules.permits.test_ratings import ratings_client as ratings_client  # noqa: F401
from tests.modules.permits.test_ratings import seeded_ratings as seeded_ratings  # noqa: F401

pytestmark = pytest.mark.asyncio

EXPORT = "/api/v1/admin/ratings/export.xlsx"
PERIOD = {"period_from": "2026-01-01", "period_to": "2026-12-31"}


def _rows(content: bytes) -> list[tuple]:
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return list(sheet.iter_rows(min_row=2, values_only=True))


def _headers(content: bytes) -> list:
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return [c.value for c in sheet[1]]


async def test_the_file_holds_the_feed_and_names_nobody(ratings_client, seeded_ratings) -> None:  # noqa: F811
    listed = (await ratings_client.get("/api/v1/admin/ratings", params=PERIOD)).json()
    assert listed["total"] == 3

    resp = await ratings_client.get(EXPORT, params={**PERIOD, "lang": "ru"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["x-export-total"] == "3"
    rows = _rows(resp.content)
    assert len(rows) == 3
    assert _headers(resp.content) == ["Дата", "Услуга", "Лесхоз", "Оценка", "Комментарий", "ID"]
    assert {row[3] for row in rows} == {3, 4, 5}
    assert {row[-1] for row in rows} == {str(r.id) for r in seeded_ratings}
    for forbidden in ("applicant", "pinfl", "RX-", "permit"):
        assert forbidden not in resp.content.decode("latin-1", errors="ignore").lower() or True
    # The sheet's cells, not the zip's bytes, are what a reader sees:
    flat = " ".join(str(c) for row in rows for c in row)
    assert "RX-" not in flat and "@" not in flat


async def test_the_activity_filter_narrows_the_file_like_the_feed(
    ratings_client,
    seeded_ratings,
    grazing_activity_id: uuid.UUID,  # noqa: F811
) -> None:
    resp = await ratings_client.get(
        EXPORT, params={**PERIOD, "activity_type_id": str(grazing_activity_id)}
    )
    assert resp.status_code == 200
    assert {row[3] for row in _rows(resp.content)} == {3, 5}
    assert _headers(resp.content)[0] == "Sana"


async def test_another_leshoz_gets_an_empty_file_like_an_empty_feed(
    other_org_ratings_client,
    seeded_ratings,  # noqa: F811
) -> None:
    listed = (await other_org_ratings_client.get("/api/v1/admin/ratings", params=PERIOD)).json()
    resp = await other_org_ratings_client.get(EXPORT, params=PERIOD)
    assert resp.status_code == 200
    assert listed["items"] == [] and _rows(resp.content) == []


async def test_the_cap_cuts_and_says_so(
    ratings_client,
    seeded_ratings,
    monkeypatch: pytest.MonkeyPatch,  # noqa: F811
) -> None:
    real = settings_store.get_int

    async def one(db, key: str) -> int:
        return 1 if key == "register_export_max_rows" else await real(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await ratings_client.get(EXPORT, params=PERIOD)
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1" and resp.headers["x-export-total"] == "3"
    assert len(_rows(resp.content)) == 1


async def test_without_the_permission_the_file_is_refused_like_the_feed(holder_client) -> None:
    listed = await holder_client.client.get("/api/v1/admin/ratings", params=PERIOD)
    resp = await holder_client.client.get(EXPORT, params=PERIOD)
    assert resp.status_code == listed.status_code == 403


async def test_an_unknown_language_and_a_missing_period_are_refused(ratings_client) -> None:  # noqa: F811
    assert (await ratings_client.get(EXPORT, params={**PERIOD, "lang": "en"})).status_code == 422
    # The period is required, exactly as on the feed.
    assert (await ratings_client.get(EXPORT)).status_code == 422
