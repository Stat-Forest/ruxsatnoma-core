"""Stage 13: the permits export is the register on paper — same scope, same
filters, readable cells, the id last.

`other_zone_hodim_client` (zoned to `other_leshoz`) is the "same caller" of
shape 1: its zone hides `issued_permit` (in `leshoz`) exactly as the list
does, while `other_zone_issued_permit` (its own leshoz's) proves the export
is not simply empty.
"""

import io

import pytest
from openpyxl import load_workbook

from app.modules.permits.models import Permit
from app.modules.permits.service import _permit_number

pytestmark = pytest.mark.asyncio


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def test_export_holds_exactly_the_rows_the_list_shows(
    other_zone_hodim_client, issued_permit: Permit, other_zone_issued_permit: Permit
):
    listed = (
        await other_zone_hodim_client.get("/api/v1/permits", params={"page_size": 100})
    ).json()
    listed_ids = {row["id"] for row in listed["items"]}
    assert str(issued_permit.id) not in listed_ids  # the zone hides the other leshoz's permit
    assert str(other_zone_issued_permit.id) in listed_ids

    resp = await other_zone_hodim_client.get("/api/v1/permits/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "№ разрешения" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids
    assert exported_ids  # non-empty: this caller has a permit to see


async def test_export_applies_the_same_filters_as_the_list(
    hodim_client, issued_permit: Permit, other_zone_issued_permit: Permit
):
    """`hodim_client` is nationwide (sees both fixture permits, plus whatever
    this shared, persistent test DB already holds — lesson), so the filter
    that proves narrowing has to pin an exact applicant, not assume "no
    matches": `applicant_id` is fresh per fixture and picks out exactly
    `issued_permit`."""
    resp = await hodim_client.get(
        "/api/v1/permits/export.xlsx",
        params={"applicant_id": str(issued_permit.applicant_id), "lang": "uz_latn"},
    )
    assert resp.status_code == 200
    rows = list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert {str(row[-1]) for row in rows} == {str(issued_permit.id)}


async def test_export_renders_labels_not_codes(hodim_client, issued_permit: Permit):
    resp = await hodim_client.get("/api/v1/permits/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[1] == "Imzolar kutilmoqda"  # the status label, not "pending_signatures"
    assert row[0] == _permit_number(
        issued_permit.series, issued_permit.number
    )  # the human number first


async def test_export_truncates_at_the_cap_and_says_so(
    hodim_client, issued_permit: Permit, monkeypatch
):
    from app.core import settings_store

    # Not just the cap: `get_current_session` also reads `session_idle_minutes`
    # through this same function on every authenticated request, so the patch
    # must fall through to the real value for any OTHER key rather than
    # asserting a single one (the plan's own template does that and would
    # break here — `deps.py::get_current_session` runs before the route body).
    real_get_int = settings_store.get_int

    async def capped(db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", capped)
    resp = await hodim_client.get("/api/v1/permits/export.xlsx")
    assert resp.headers["x-export-truncated"] == (
        "true" if int(resp.headers["x-export-total"]) > 1 else "false"
    )
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_export_is_empty_not_403_for_a_caller_with_no_scope(other_applicant_client):
    list_resp = await other_applicant_client.client.get("/api/v1/permits")
    assert list_resp.status_code == 200
    assert list_resp.json()["items"] == []

    resp = await other_applicant_client.client.get("/api/v1/permits/export.xlsx")
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_export_rejects_an_unknown_language(hodim_client):
    assert (
        await hodim_client.get("/api/v1/permits/export.xlsx", params={"lang": "en"})
    ).status_code == 422
