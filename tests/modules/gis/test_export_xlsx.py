"""Stage 13: the contours/imports exports are the lists on paper — same
filters, same zone scoping, readable cells, the id last, NO geometry."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.models import MediaFile
from app.modules.admin.models import Organization
from app.modules.gis.models import Contour, GisLayer
from tests.modules.gis.conftest import make_contour, make_version, random_box_wkt

pytestmark = pytest.mark.asyncio


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def _published_contour(
    db: AsyncSession,
    layer: GisLayer,
    org: Organization,
    approval_doc: MediaFile,
    **contour_over,
) -> Contour:
    """A fresh contour with one published version — the plain shape most of
    these tests need, built from the module's own `make_contour`/
    `make_version` helpers rather than the package's `published_contour`
    fixture (which hands back the `ContourVersion`, not the `Contour` a
    `.number` assertion needs)."""
    contour = await make_contour(db, layer, org, **contour_over)
    await make_version(
        db,
        contour.id,
        random_box_wkt(),
        status="published",
        approval_doc_id=approval_doc.id,
        published_at=func.now(),
    )
    await db.flush()
    return contour


# --- contours --------------------------------------------------------------


async def test_contours_export_holds_exactly_the_rows_the_list_shows(
    applicant_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    approval_doc: MediaFile,
) -> None:
    await _published_contour(db, contours_layer, leshoz, approval_doc)

    listed = await applicant_client.get(f"/api/v1/gis/contours?organization_id={leshoz.id}")
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert listed_ids

    resp = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx",
        params={"organization_id": str(leshoz.id), "lang": "ru"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Номер" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_contours_export_applies_the_same_bbox_filter_as_the_list(
    applicant_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    approval_doc: MediaFile,
) -> None:
    await _published_contour(db, contours_layer, leshoz, approval_doc)

    inside = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx",
        params={"organization_id": str(leshoz.id), "bbox": "0,0,40,30"},
    )
    outside = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx",
        params={"organization_id": str(leshoz.id), "bbox": "60.0,41.4,60.1,41.5"},
    )
    assert inside.status_code == 200, inside.text
    assert outside.status_code == 200, outside.text
    assert len(list(_sheet(inside.content).iter_rows(min_row=2, values_only=True))) >= 1
    assert list(_sheet(outside.content).iter_rows(min_row=2, values_only=True)) == []


async def test_contours_export_rejects_a_malformed_bbox_like_the_list(
    applicant_client: AsyncClient,
) -> None:
    resp = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx", params={"bbox": "nonsense"}
    )
    assert resp.status_code == 422


async def test_contours_export_renders_labels_not_codes(
    applicant_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    approval_doc: MediaFile,
) -> None:
    contour = await _published_contour(db, contours_layer, leshoz, approval_doc)

    resp = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx",
        params={"organization_id": str(leshoz.id), "lang": "uz_latn"},
    )
    rows = {row[0]: row for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)}
    assert contour.number in rows  # the human number first
    row = rows[contour.number]
    assert row[3] == "Kontur"  # the kind label, not "contour"
    assert row[8] == "Yoʻq"  # over_allocated label, not a bare Python bool


async def test_contours_export_truncates_at_the_cap_and_says_so(
    applicant_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    approval_doc: MediaFile,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await _published_contour(db, contours_layer, leshoz, approval_doc)
    await _published_contour(db, contours_layer, leshoz, approval_doc)

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db_arg, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_arg, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await applicant_client.get(
        "/api/v1/gis/contours/export.xlsx", params={"organization_id": str(leshoz.id)}
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["x-export-total"] == "2"
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_contours_export_is_empty_for_a_caller_with_no_matching_zone(
    region_scoped_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    approval_doc: MediaFile,
) -> None:
    """`region_scoped_client` is zoned to a region `leshoz` does not belong
    to (`leshoz` sets no region at all) — the same empty-not-403 shape
    `test_the_list_is_filtered_for_a_region_scoped_actor` proves for the
    list."""
    await _published_contour(db, contours_layer, leshoz, approval_doc)

    resp = await region_scoped_client.get("/api/v1/gis/contours/export.xlsx")
    assert resp.status_code == 200, resp.text
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_contours_export_rejects_an_unknown_language(applicant_client: AsyncClient) -> None:
    resp = await applicant_client.get("/api/v1/gis/contours/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- imports -----------------------------------------------------------------


async def test_imports_export_holds_exactly_the_rows_the_list_shows(
    gis_client: AsyncClient, pending_import
) -> None:
    listed = await gis_client.get("/api/v1/gis/imports")
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert str(pending_import.id) in listed_ids

    resp = await gis_client.get("/api/v1/gis/imports/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Создан" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_imports_export_applies_the_same_status_filter_as_the_list(
    rahbar_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    gis_user,
) -> None:
    from tests.modules.gis.conftest import make_import

    review_batch = await make_import(
        db, layer=contours_layer, org=leshoz, started_by=gis_user, data=b"{}"
    )
    review_batch.status = "review"
    await db.flush()
    await db.commit()
    await make_import(db, layer=contours_layer, org=leshoz, started_by=gis_user, data=b"{}")

    resp = await rahbar_client.get("/api/v1/gis/imports/export.xlsx", params={"status": "review"})
    assert resp.status_code == 200, resp.text
    ids = {str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)}
    assert str(review_batch.id) in ids


async def test_imports_export_renders_labels_and_the_filename_not_the_file_id(
    gis_client: AsyncClient, pending_import
) -> None:
    resp = await gis_client.get("/api/v1/gis/imports/export.xlsx", params={"lang": "uz_latn"})
    row = next(
        r
        for r in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
        if str(r[-1]) == str(pending_import.id)
    )
    assert row[1] == "Navbatda"  # the status label, not "pending"
    assert row[5] == "import.geojson"  # the filename, not a media_files uuid


async def test_imports_export_truncates_at_the_cap_and_says_so(
    gis_client: AsyncClient,
    db: AsyncSession,
    contours_layer: GisLayer,
    leshoz: Organization,
    gis_user,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from tests.modules.gis.conftest import make_import

    await make_import(db, layer=contours_layer, org=leshoz, started_by=gis_user, data=b"{}")
    await make_import(db, layer=contours_layer, org=leshoz, started_by=gis_user, data=b"{}")

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db_arg, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_arg, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await gis_client.get("/api/v1/gis/imports/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_imports_export_is_empty_for_a_caller_zoned_elsewhere(
    org_scoped_rahbar_client: AsyncClient, pending_import
) -> None:
    """`pending_import` sits under `leshoz`; `org_scoped_rahbar_client` is
    zoned to `other_leshoz` — the same shape `test_the_import_list_is_zone_scoped`
    proves for the JSON list."""
    resp = await org_scoped_rahbar_client.get("/api/v1/gis/imports/export.xlsx")
    assert resp.status_code == 200, resp.text
    ids = {str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)}
    assert str(pending_import.id) not in ids


async def test_imports_export_refuses_a_caller_with_neither_manage_nor_approve(
    applicant_client: AsyncClient,
) -> None:
    resp = await applicant_client.get("/api/v1/gis/imports/export.xlsx")
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_imports_export_rejects_an_unknown_language(gis_client: AsyncClient) -> None:
    resp = await gis_client.get("/api/v1/gis/imports/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
