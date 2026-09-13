"""`GET /applications/export.xlsx` is the applications register on paper —
the same scope, the same filters, readable cells, the id last (stage 13,
ruling #204), with the columns the Agency's PM asked for on 2026-09-13:
region, district, phone, benefit, the permit's issue date and term, the
calculated and the paid amount, the executor's conclusion, and the status
both as the customer's six groups and as our exact one.

The route is served by `reports.applications_register` — a level-5 reader,
because half of those columns live in `permits`, `payments` and `norms`,
which `applications` (level 3) may not read. The tests stay in THIS package
because the fixtures that file an application live here.

Every scope-shaped test compares the FILE against the LIST for the same
caller, because the one defect this route could carry that matters is a
disagreement between the two (ruling #98: an export that widens a zone; or,
the direction this project's defects actually take, one that silently
narrows it).
"""

import io
import uuid
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import settings_store, xlsx
from app.modules.admin.models import District, Organization, Region
from app.modules.applications import service as applications_service
from app.modules.applications.models import APPLICATION_STATUSES
from app.modules.auth.models import Applicant
from app.modules.reports import applications_register as register

pytestmark = pytest.mark.asyncio

EXPORT = "/api/v1/applications/export.xlsx"


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


def _data_rows(content: bytes) -> list[tuple]:
    return list(_sheet(content).iter_rows(min_row=2, values_only=True))


def _by_header(content: bytes, application_id: str) -> dict[str, Any]:
    headers = [str(c.value) for c in _sheet(content)[1]]
    for row in _data_rows(content):
        if row[-1] == application_id:
            return dict(zip(headers, row, strict=True))
    raise AssertionError(f"{application_id} is not in the file")


async def test_the_export_holds_exactly_the_rows_the_list_shows(
    submitted_application: str, published_contour, hodim_client, other_zone_hodim_client
) -> None:
    listed = await hodim_client.get("/api/v1/applications", params={"page_size": 100})
    listed_ids = {row["id"] for row in listed.json()["items"]}
    # A NON-EMPTY scope, or nothing below proves anything.
    assert submitted_application in listed_ids

    resp = await hodim_client.get(EXPORT, params={"lang": "ru"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    assert resp.headers["x-export-total"] == str(len(listed_ids))
    assert resp.headers["content-disposition"].startswith('attachment; filename="arizalar-')

    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "№" and headers[-1] == "ID"
    rows = _data_rows(resp.content)
    assert {row[-1] for row in rows} == listed_ids
    # The serial number counts the rows of THIS file, 1..n, whatever the ids.
    assert [row[0] for row in rows] == list(range(1, len(rows) + 1))

    # The other leshoz's hodim: the same contour filter empties both the list
    # and the file — the zone hides the row in both places, never in one.
    theirs = await other_zone_hodim_client.get(
        EXPORT, params={"contour_id": str(published_contour.id)}
    )
    assert theirs.status_code == 200
    assert _data_rows(theirs.content) == []
    assert theirs.headers["x-export-total"] == "0"


async def test_the_export_applies_the_same_filters_as_the_list(
    submitted_application: str, published_contour, hodim_client
) -> None:
    narrowed = await hodim_client.get(
        EXPORT, params={"contour_id": str(published_contour.id), "lang": "uz_latn"}
    )
    assert narrowed.status_code == 200
    assert [row[-1] for row in _data_rows(narrowed.content)] == [submitted_application]

    none = await hodim_client.get(EXPORT, params={"status": "REJECTED"})
    assert none.status_code == 200
    assert _data_rows(none.content) == []

    typo = await hodim_client.get(EXPORT, params={"status": "SUBMITED"})
    assert typo.status_code == 422  # the same literal the list validates


async def test_the_export_renders_names_labels_and_places_not_codes(
    db: AsyncSession,
    submitted_application: str,
    published_contour,
    applicant: Applicant,
    leshoz: Organization,
    hodim_client,
) -> None:
    """Every id the row carries is resolved: the applicant to a name and a
    phone, the leshoz to its name AND its region and district, the activity
    to a name, the unit to a word, the status to the customer's group and
    to our exact label."""
    region = (await db.execute(select(Region).where(Region.code == "fergana"))).scalar_one()
    # Districts come from the seed CLI, not a migration, so the test database
    # has none: get-or-create one by its unique code (the DB is shared and
    # persistent, so a second run finds the first run's row).
    district = (
        await db.execute(select(District).where(District.code == "test-export-district"))
    ).scalar_one_or_none()
    if district is None:
        district = District(
            code="test-export-district",
            name={"uz_latn": "Eksport tumani", "ru": "Экспортный район"},
            region_id=region.id,
        )
        db.add(district)
        await db.flush()
    leshoz.region_id = region.id
    leshoz.district_id = district.id
    applicant.phone = "+998901234567"
    await db.commit()

    resp = await hodim_client.get(
        EXPORT, params={"contour_id": str(published_contour.id), "lang": "uz_latn"}
    )
    assert resp.status_code == 200, resp.text
    row = _by_header(resp.content, submitted_application)

    assert row["№"] == 1
    assert isinstance(row["Ariza raqami"], str) and row["Ariza raqami"].startswith("RX-")
    assert row["Viloyat"] == xlsx.localized(region.name, "uz_latn")
    assert row["Tuman"] == xlsx.localized(district.name, "uz_latn")
    assert row["Oʻrmon xoʻjaligi"] == xlsx.localized(leshoz.name, "uz_latn")
    assert row["Kontur"] == published_contour.number  # a number, not a UUID
    assert row["F.I.Sh."] == applicant.name
    assert row["Telefon"] == "+998901234567"
    assert row["Ruxsatnoma turi"]  # resolved to a name
    assert row["Imtiyoz turi"] is None  # no benefit claimed — an empty cell
    assert row["Miqdor"] == 40  # the filing's 40 head, as a NUMBER
    assert row["Birlik"] == "bosh"  # the unit as a word, beside it
    assert row["Ariza holati"] == "Yangi ariza"  # the customer's group
    assert row["Holati"] == "Yuborilgan"  # our exact label, not SUBMITTED
    assert row["Ruxsatnoma berilgan sana"] is None  # no permit yet
    assert row["Ruxsatnoma muddati: dan"] is None
    assert row["Hisoblangan summa"] is not None  # the filing priced itself
    assert row["Toʻlangan summa"] is None  # nothing paid
    assert row["Xulosa"] is None  # nobody has written one
    assert row["ID"] == submitted_application

    in_russian = await hodim_client.get(
        EXPORT, params={"contour_id": str(published_contour.id), "lang": "ru"}
    )
    row_ru = _by_header(in_russian.content, submitted_application)
    assert row_ru["Статус"] == "Подана"
    assert row_ru["Статус заявки"] == "Новая заявка"
    assert row_ru["Ед. изм."] == "голова"
    assert row_ru["Область"] == xlsx.localized(region.name, "ru")


async def test_the_export_carries_the_review_columns(
    db: AsyncSession, application_in_review: str, hodim_client, executor_head_client
) -> None:
    """After the executor's conclusion and the head's approval the row shows
    the conclusion's text, the calculated amount and the customer's
    «reviewed» group — while the paid amount stays empty, because nothing
    has been paid (an invoice is not a payment)."""
    from tests.modules.applications.test_decision import _decide

    written = await hodim_client.post(
        f"/api/v1/applications/{application_in_review}/conclusion",
        json={"text": "Комплект полный", "kind": "executor", "recommendation": "approve"},
    )
    assert written.status_code == 201, written.text
    decided = await _decide(executor_head_client, application_in_review, "approve")
    assert decided.status_code == 200, decided.text

    calculation = await applications_service.current_calculation(
        db, uuid.UUID(application_in_review)
    )
    assert calculation is not None

    resp = await hodim_client.get(EXPORT, params={"lang": "uz_latn"})
    assert resp.status_code == 200, resp.text
    row = _by_header(resp.content, application_in_review)
    assert row["Xulosa"] == "Комплект полный"
    assert row["Hisoblangan summa"] == calculation.amount
    assert row["Toʻlangan summa"] is None
    assert row["Holati"] == "Hisob-faktura yuborilgan"  # INVOICED, by the bus
    assert row["Ariza holati"] == "Koʻrib chiqilgan"


async def test_every_status_falls_into_exactly_one_customer_group() -> None:
    """The six groups are the customer's vocabulary over our thirteen
    statuses; a status without a group would render as its own code, and a
    group without a label in one language would refuse to render at all."""
    for status in APPLICATION_STATUSES:
        group = register.STATUS_GROUPS[status]
        assert set(register.GROUP_LABELS[group]) == set(xlsx.LANGS), (status, group)
    assert set(register.STATUS_GROUPS) == set(APPLICATION_STATUSES)
    for status in APPLICATION_STATUSES:
        assert set(register.STATUS_LABELS[status]) == set(xlsx.LANGS), status


async def test_the_export_truncates_at_the_cap_and_says_so(
    submitted_application: str, hodim_client, monkeypatch: pytest.MonkeyPatch
) -> None:
    real_get_int = settings_store.get_int

    async def one(db, key: str) -> int:
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)

    listed = await hodim_client.get("/api/v1/applications", params={"page_size": 100})
    total = listed.json()["total"]
    assert total >= 1

    resp = await hodim_client.get(EXPORT)
    assert resp.status_code == 200
    assert resp.headers["x-export-total"] == str(total)
    assert resp.headers["x-export-rows"] == "1"
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(_data_rows(resp.content)) == 1


async def test_the_file_and_the_list_agree_for_a_republic_wide_reader(
    submitted_application: str, prosecutor_client
) -> None:
    """The prosecutor's zone is the whole republic: the file holds the same
    ids the list does — neither wider (ruling #98) nor narrower.

    The list is read to its END, not as one page of 100: nothing rolls a test
    back (`tests/conftest.py`), so this worker's database holds every
    application the module's earlier tests filed, and once they pass a hundred
    a single page is a strict subset of the file (CI, 2026-09-13)."""
    listed_ids: set[str] = set()
    page = 1
    while True:
        listed = await prosecutor_client.get(
            "/api/v1/applications", params={"page": page, "page_size": 100}
        )
        assert listed.status_code == 200, listed.text
        body = listed.json()
        listed_ids |= {row["id"] for row in body["items"]}
        if page * 100 >= body["total"]:
            break
        page += 1
    resp = await prosecutor_client.get(EXPORT)
    assert resp.status_code == 200
    assert {row[-1] for row in _data_rows(resp.content)} == listed_ids


async def test_an_unknown_language_is_refused(hodim_client) -> None:
    assert (await hodim_client.get(EXPORT, params={"lang": "en"})).status_code == 422


async def test_the_export_route_is_not_shadowed_by_the_card_route(hodim_client) -> None:
    """`export.xlsx` is not a UUID: mounted after `/{application_id}` it would
    be a 422 from the id parser rather than this route — and since the route
    now lives in another module's router, the ORDER `main.py` mounts the two
    routers in is what keeps it so."""
    resp = await hodim_client.get(EXPORT)
    assert resp.status_code == 200, resp.text
    assert (await hodim_client.get(f"/api/v1/applications/{uuid.uuid4()}")).status_code == 404
