"""Stage 13: `GET /search/export.xlsx` — the plain register export beside
the prosecutor's watermarked `POST /search/exports` (untouched). The export
is the list on paper — same zone, same filters, readable cells, the id
last."""

import io

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.admin.models import Organization
from app.modules.search.permissions import SEARCH_USE
from tests.modules.search.conftest import _client_for, _client_with_role, make_application


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def test_export_holds_exactly_the_rows_the_list_shows(
    db: AsyncSession, leshoz: Organization, other_leshoz: Organization
):
    mine = await make_application(db, org=leshoz, status="CLOSED", applicant_name="A A")
    theirs = await make_application(db, org=other_leshoz, status="CLOSED", applicant_name="B B")
    await db.commit()

    async for client in _client_for(db, SEARCH_USE, organization_id=leshoz.id):
        listed = await client.get(
            "/api/v1/search", params={"kind": "applications", "page_size": 100}
        )
        listed_ids = {row["id"] for row in listed.json()["items"]}
        assert str(theirs.id) not in listed_ids  # the zone hides it from the screen

        resp = await client.get(
            "/api/v1/search/export.xlsx", params={"kind": "applications", "lang": "ru"}
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert resp.headers["x-export-truncated"] == "false"
        sheet = _sheet(resp.content)
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "Номер" and headers[-1] == "ID"
        exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert str(mine.id) in exported_ids
        assert exported_ids == listed_ids


async def test_export_applies_the_same_filters_as_the_list(db: AsyncSession, leshoz: Organization):
    closed = await make_application(db, org=leshoz, status="CLOSED", applicant_name="A A")
    cancelled = await make_application(db, org=leshoz, status="CANCELLED", applicant_name="B B")
    await db.commit()

    async for client in _client_for(db, SEARCH_USE, organization_id=leshoz.id):
        resp = await client.get(
            "/api/v1/search/export.xlsx",
            params={"kind": "applications", "status": "CLOSED", "lang": "uz_latn"},
        )
        assert resp.status_code == 200
        exported_ids = {
            str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
        }
        assert str(closed.id) in exported_ids
        assert str(cancelled.id) not in exported_ids


async def test_export_renders_labels_not_codes_and_the_number_first(
    db: AsyncSession, leshoz: Organization
):
    app = await make_application(db, org=leshoz, status="CLOSED", applicant_name="A A")
    await db.commit()

    async for client in _client_for(db, SEARCH_USE, organization_id=leshoz.id):
        resp = await client.get(
            "/api/v1/search/export.xlsx", params={"kind": "applications", "lang": "uz_latn"}
        )
        row = next(
            r
            for r in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
            if r[-1] == str(app.id)
        )
        assert row[0] == app.number  # the human number first
        assert row[1] == "Ariza"  # kind label, not the raw code
        assert row[4] == "Yopilgan"  # the status LABEL, not CLOSED


async def test_export_truncates_at_the_cap_and_says_so(
    db: AsyncSession, leshoz: Organization, monkeypatch
):
    for i in range(3):
        await make_application(db, org=leshoz, status="CLOSED", applicant_name=f"Cap {i}")
    await db.commit()

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def capped(db_, key):
        # `auth.deps.get_current_session` reads `session_idle_minutes`
        # through this SAME function on every request — only the export's
        # own cap key is overridden here, every other key falls through to
        # the real function or authentication itself breaks.
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_, key)

    monkeypatch.setattr(settings_store, "get_int", capped)

    async for client in _client_for(db, SEARCH_USE, organization_id=leshoz.id):
        resp = await client.get("/api/v1/search/export.xlsx", params={"kind": "applications"})
        assert resp.status_code == 200
        total = int(resp.headers["x-export-total"])
        assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
        assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1
        assert resp.headers["x-export-rows"] == str(min(total, 1))


async def test_export_requires_the_permission(db: AsyncSession, leshoz: Organization):
    async for client in _client_with_role(db, "inspector"):  # search.use NOT granted
        resp = await client.get("/api/v1/search/export.xlsx", params={"kind": "applications"})
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_export_rejects_an_unknown_language(db: AsyncSession, leshoz: Organization):
    async for client in _client_for(db, SEARCH_USE, organization_id=leshoz.id):
        resp = await client.get(
            "/api/v1/search/export.xlsx", params={"kind": "applications", "lang": "en"}
        )
        assert resp.status_code == 422
