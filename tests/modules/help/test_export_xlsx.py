"""Stage 13: `/help/tickets/export.xlsx` and `/admin/help/faq/export.xlsx`
are their lists on paper — same scope/permission, same filters, readable
cells, the id last."""

import io

import pytest
from openpyxl import load_workbook

from app.main import create_app
from tests.conftest import make_client
from tests.modules.auth.test_sessions import make_session, make_user
from tests.modules.help.conftest import auth_client, faq_manager

pytestmark = pytest.mark.asyncio

API = "/api/v1"


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


def _ids(sheet) -> set[str]:
    """The last column of every data row, as strings — `iter_rows(values_only
    =True)` types a cell as a broad union openpyxl itself does not guarantee
    hashable, so every id is coerced through `str()` before it enters a set
    (every id column ever holds one already, via `xlsx.id_column`)."""
    return {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}


def _row_by_id(sheet, id_: str):
    """The FAQ catalogue has no per-caller scope (`help.faq.manage` sees
    every row, published or not), so the shared, persistent test DB can
    hold rows from earlier runs — `next(iter_rows(...))` would read
    whichever one happens to sort first, not necessarily this fixture's own
    (lesson: filter on something fresh per fixture, never assume order)."""
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[-1]) == id_:
            return row
    raise AssertionError(f"no row with id {id_!r} in the sheet")


async def _signed_in(db, **overrides):
    user = await make_user(db, **overrides)
    _, token, csrf = await make_session(db, user)
    await db.commit()
    return user, token, csrf


def _cap_one_but_keep_session_idle_minutes(monkeypatch):
    """`auth.deps.get_current_session` reads `session_idle_minutes` through
    the SAME `settings_store.get_int` on every request — the patch must fall
    through to the real function for every key but ours (lesson)."""
    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)


# --- Tickets -----------------------------------------------------------


async def test_tickets_export_holds_exactly_the_rows_the_list_shows(db):
    owner, token, csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/help/tickets", json={"subject": "Cannot pay", "body": "Help"}
        )
        assert created.status_code == 201, created.text
        ticket_id = created.json()["id"]

        listed = await client.get(f"{API}/help/tickets", params={"page_size": 100})
        listed_ids = {row["id"] for row in listed.json()["items"]}
        assert ticket_id in listed_ids

        resp = await client.get(f"{API}/help/tickets/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Номер" and headers[-1] == "ID"
    exported_ids = _ids(sheet)
    assert exported_ids == listed_ids
    assert ticket_id in exported_ids


async def test_tickets_export_applies_the_same_status_filter_as_the_list(db):
    owner, token, csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/help/tickets", json={"subject": "Open one", "body": "b"}
        )
        assert created.status_code == 201, created.text

        resp = await client.get(
            f"{API}/help/tickets/export.xlsx", params={"status": "closed", "lang": "uz_latn"}
        )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_tickets_export_renders_labels_not_codes(db):
    owner, token, csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(f"{API}/help/tickets", json={"subject": "Zebra", "body": "b"})
        assert created.status_code == 201, created.text

        resp = await client.get(f"{API}/help/tickets/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == created.json()["number"]  # the human number comes first
    assert row[1] == "Zebra"
    assert row[2] == "Yangi"  # status label ("new"), not the raw code


async def test_tickets_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    owner, token, csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(f"{API}/help/tickets", json={"subject": "Capped", "body": "b"})
        assert created.status_code == 201, created.text

        _cap_one_but_keep_session_idle_minutes(monkeypatch)
        resp = await client.get(f"{API}/help/tickets/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_tickets_export_is_empty_not_403_for_a_caller_with_no_scope(db):
    owner, owner_token, owner_csrf = await _signed_in(db)
    stranger, stranger_token, stranger_csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, owner_token, owner_csrf)
        created = await client.post(
            f"{API}/help/tickets", json={"subject": "Owner's own", "body": "b"}
        )
        assert created.status_code == 201, created.text

        auth_client(client, stranger_token, stranger_csrf)
        resp = await client.get(f"{API}/help/tickets/export.xlsx")
    assert resp.status_code == 200
    # `stranger` has never opened or been assigned a ticket — an empty 200,
    # the same shape the list gives a caller in their position, never a 403.
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_tickets_export_rejects_an_unknown_language(db):
    _, token, csrf = await _signed_in(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/help/tickets/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- FAQ (admin) -----------------------------------------------------------


async def test_faq_export_holds_exactly_the_rows_the_list_shows(db):
    _, token, csrf = await faq_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/admin/help/faq",
            json={
                "question": {"uz_cyrl": "Саволми?", "uz_latn": "Savolmi?"},
                "answer": {"uz_cyrl": "Жавоб.", "uz_latn": "Javob."},
                "sort_order": 3,
            },
        )
        assert created.status_code == 201, created.text
        faq_id = created.json()["id"]

        listed = await client.get(f"{API}/admin/help/faq")
        listed_ids = {row["id"] for row in listed.json()}
        assert faq_id in listed_ids

        resp = await client.get(f"{API}/admin/help/faq/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Вопрос" and headers[-1] == "ID"
    exported_ids = _ids(sheet)
    assert exported_ids == listed_ids
    assert faq_id in exported_ids


async def test_faq_export_applies_the_same_status_filter_as_the_list(db):
    _, token, csrf = await faq_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/admin/help/faq",
            json={
                "question": {"uz_cyrl": "Q?", "uz_latn": "Q?"},
                "answer": {"uz_cyrl": "A.", "uz_latn": "A."},
            },
        )
        assert created.status_code == 201, created.text  # status defaults to "draft"

        resp = await client.get(
            f"{API}/admin/help/faq/export.xlsx", params={"status": "archived", "lang": "uz_latn"}
        )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_faq_export_renders_labels_not_codes(db):
    _, token, csrf = await faq_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/admin/help/faq",
            json={
                "question": {"uz_cyrl": "Нархими?", "uz_latn": "Narxi qancha?"},
                "answer": {"uz_cyrl": "Ж.", "uz_latn": "J."},
                "category": "billing",
            },
        )
        assert created.status_code == 201, created.text
        faq_id = created.json()["id"]
        published = await client.patch(
            f"{API}/admin/help/faq/{faq_id}", json={"status": "published"}
        )
        assert published.status_code == 200, published.text

        resp = await client.get(f"{API}/admin/help/faq/export.xlsx", params={"lang": "uz_latn"})
    row = _row_by_id(_sheet(resp.content), faq_id)
    assert row[0] == "Narxi qancha?"  # localized question, not the raw dict
    assert row[1] == "billing"
    assert row[2] == "Chop etilgan"  # status label ("published"), not the raw code


async def test_faq_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await faq_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            f"{API}/admin/help/faq",
            json={
                "question": {"uz_cyrl": "Q?", "uz_latn": "Q?"},
                "answer": {"uz_cyrl": "A.", "uz_latn": "A."},
            },
        )
        assert created.status_code == 201, created.text

        _cap_one_but_keep_session_idle_minutes(monkeypatch)
        resp = await client.get(f"{API}/admin/help/faq/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_faq_export_requires_the_same_permission_as_the_list(db):
    stranger = await make_user(db)
    _, token, csrf = await make_session(db, stranger)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        list_resp = await client.get(f"{API}/admin/help/faq")
        export_resp = await client.get(f"{API}/admin/help/faq/export.xlsx")
    assert list_resp.status_code == 403
    assert export_resp.status_code == 403
    assert export_resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_faq_export_rejects_an_unknown_language(db):
    _, token, csrf = await faq_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/help/faq/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
