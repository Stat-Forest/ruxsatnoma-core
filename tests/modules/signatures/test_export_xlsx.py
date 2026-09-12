"""Stage 13: `GET /certificates/export.xlsx` and `GET /signatures/export.xlsx`
— the caller's own bound certificates and one object's signature list on
paper. Neither is wired to an Excel button by this track (out of its screen
list): `src/pages/admin/profile/certificates/CertificatesSection.tsx` and
`src/pages/permits/PermitSignaturesPanel.tsx` DO read these routes today,
but each shows a small, single-owner/single-object list ("a handful of
keys at most", that file's own comment) rather than a register — track
report flags the correction. Fixtures reused from `test_api.py` (pre-flight
ruling P4: this file adds its own where `test_api.py`'s do not fit, rather
than growing the shared `conftest.py`)."""

import io
import uuid

from openpyxl import load_workbook

from app.modules.auth.models import User
from app.modules.signatures import service
from app.modules.signatures.models import Signature
from tests.modules.signatures.test_api import (
    DOC,
    _pkcs7,
)
from tests.modules.signatures.test_api import (
    a_signature as a_signature,
)
from tests.modules.signatures.test_api import (
    bound_cert_of_a as bound_cert_of_a,
)
from tests.modules.signatures.test_api import (
    client_a as client_a,
)
from tests.modules.signatures.test_api import (
    client_b as client_b,
)
from tests.modules.signatures.test_api import (
    client_oversight as client_oversight,
)
from tests.modules.signatures.test_api import (
    user_a as user_a,
)
from tests.modules.signatures.test_api import (
    user_b as user_b,
)


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


# --- /certificates/export.xlsx ---------------------------------------------


async def test_certificates_export_holds_exactly_the_rows_the_list_shows(
    client_a, client_b, bound_cert_of_a
):
    mine = await client_a.get("/api/v1/certificates")
    assert [c["id"] for c in mine.json()["items"]] == [str(bound_cert_of_a)]

    resp = await client_a.get("/api/v1/certificates/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Субъект" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == {str(bound_cert_of_a)}

    theirs = await client_b.get("/api/v1/certificates/export.xlsx")
    assert theirs.status_code == 200
    assert list(_sheet(theirs.content).iter_rows(min_row=2, values_only=True)) == []


async def test_certificates_export_renders_the_owner(client_a, bound_cert_of_a, user_a: User):
    resp = await client_a.get("/api/v1/certificates/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[5] == user_a.full_name  # the resolved user name, not a bare id


async def test_certificates_export_truncates_at_the_cap_and_says_so(
    client_a, user_a: User, monkeypatch
):
    for _ in range(2):
        pkcs7 = _pkcs7(b"another-challenge-" + uuid.uuid4().bytes, user_a.pinfl)
        resp = await client_a.post("/api/v1/certificates", json={"pkcs7": pkcs7})
        assert resp.status_code == 201, resp.text

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def capped(db_, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_, key)

    monkeypatch.setattr(settings_store, "get_int", capped)

    resp = await client_a.get("/api/v1/certificates/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1
    assert resp.headers["x-export-rows"] == str(min(total, 1))


async def test_certificates_export_is_empty_not_an_error_for_a_caller_with_none(client_b):
    """No permission gate on `/certificates` at all — every authenticated
    caller gets their OWN (possibly empty) list; this is the list's own
    shape for "no rows", mirrored here (brief's shape 5)."""
    resp = await client_b.get("/api/v1/certificates/export.xlsx")
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_certificates_export_rejects_an_unknown_language(client_a):
    resp = await client_a.get("/api/v1/certificates/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- /signatures/export.xlsx -----------------------------------------------


async def test_signatures_export_holds_exactly_the_rows_the_list_shows(
    client_a, a_signature: Signature, user_a: User
):
    listed = await client_a.get(
        f"/api/v1/signatures?object_type=permit&object_id={a_signature.object_id}"
    )
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert str(a_signature.id) in listed_ids

    resp = await client_a.get(
        "/api/v1/signatures/export.xlsx",
        params={"object_type": "permit", "object_id": str(a_signature.object_id), "lang": "ru"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Тип объекта" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_signatures_export_renders_labels_and_the_signer(
    client_a, a_signature: Signature, user_a: User
):
    resp = await client_a.get(
        "/api/v1/signatures/export.xlsx",
        params={
            "object_type": "permit",
            "object_id": str(a_signature.object_id),
            "lang": "uz_latn",
        },
    )
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[3] == "ERI"  # kind LABEL, not "eri"
    assert row[4] == user_a.full_name  # the resolved signer name
    assert row[6] == "Haqiqiy"  # verification_status LABEL, not "valid"


async def test_signatures_export_truncates_at_the_cap_and_says_so(
    client_a, user_a: User, a_signature: Signature, db, monkeypatch
):
    # A SECOND valid signature against the SAME object, under a different
    # purpose (`uq_signatures_valid_purpose` allows one valid signature per
    # purpose, not per object) -- so this object now carries 2 rows to cap.
    await service.sign(
        db,
        object_type="permit",
        object_id=a_signature.object_id,
        purpose="permit_recipient",
        document=DOC,
        pkcs7=_pkcs7(DOC, user_a.pinfl),
        user=user_a,
    )
    await db.commit()

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def capped(db_, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_, key)

    monkeypatch.setattr(settings_store, "get_int", capped)

    resp = await client_a.get(
        "/api/v1/signatures/export.xlsx",
        params={"object_type": "permit", "object_id": str(a_signature.object_id)},
    )
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_signatures_export_requires_ownership_or_view_any(client_b, a_signature: Signature):
    """The SAME status the list gives a non-owner without `view_any`
    (brief's shape 5)."""
    resp = await client_b.get(
        "/api/v1/signatures/export.xlsx",
        params={"object_type": "permit", "object_id": str(a_signature.object_id)},
    )
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_signatures_export_rejects_an_unknown_language(client_a, a_signature: Signature):
    resp = await client_a.get(
        "/api/v1/signatures/export.xlsx",
        params={
            "object_type": "permit",
            "object_id": str(a_signature.object_id),
            "lang": "en",
        },
    )
    assert resp.status_code == 422
