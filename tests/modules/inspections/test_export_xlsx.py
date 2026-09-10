"""Stage 13: `/inspections/{tasks,acts,cases}/export.xlsx` are their lists on
paper — same scope, same filters, readable cells, the id last."""

import io

import pytest
from openpyxl import load_workbook

from app.modules.applications.models import Application
from app.modules.inspections import repo as inspections_repo
from app.modules.inspections import service as inspections_service
from app.modules.integrations.adapters.eimzo import encode_mock_signature

pytestmark = pytest.mark.asyncio

API = "/api/v1/inspections"


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


def _ids(sheet) -> set[str]:
    """The last column of every data row, as strings — `iter_rows(values_only
    =True)` types a cell as a broad union openpyxl itself does not guarantee
    hashable, so every id is coerced through `str()` before it enters a set
    (every id column ever holds one already, via `xlsx.id_column`)."""
    return {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}


async def _create_task(client, application: Application, inspector) -> str:
    r = await client.post(
        f"{API}/tasks",
        json={
            "kind": "permit_inspection",
            "application_id": str(application.id),
            "assigned_to": str(inspector.id),
        },
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


async def _create_act(
    client, application: Application, checklist_id, *, result: str = "compliant"
) -> str:
    answers = {"activity_matches": True, "within_contour": result != "violation"}
    r = await client.post(
        f"{API}/acts",
        json={
            "application_id": str(application.id),
            "occurred_at": "2027-06-01T10:00:00Z",
            "checklist_id": str(checklist_id),
            "answers": answers,
            "result": result,
        },
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


async def _open_case(db, inspector_client, inspector, application, checklist_id, vt_id) -> str:
    """A signed `violation` act opens a case automatically (`service.sign_act`)
    — never assigned by hand (lesson: build a fixture's precondition through
    the real transition)."""
    import uuid as uuid_mod

    act_id = await _create_act(inspector_client, application, checklist_id, result="violation")
    act = await inspections_repo.get_act(db, uuid_mod.UUID(act_id))
    assert act is not None
    pkcs7 = encode_mock_signature(
        document=inspections_service._act_package_bytes(act),
        serial=f"SN-{inspector.pinfl}",
        issuer="ISS-1",
        pinfl=inspector.pinfl,
    )
    signed = await inspector_client.post(
        f"{API}/acts/{act_id}/sign", json={"pkcs7": pkcs7, "violation_type_item_id": str(vt_id)}
    )
    assert signed.status_code == 200, signed.text
    case = await inspections_repo.case_for_act(db, act.id)
    assert case is not None
    return str(case.id)


# --- Tasks -----------------------------------------------------------------


async def test_tasks_export_holds_exactly_the_rows_the_list_shows(
    executor_head_client, application: Application, inspector
) -> None:
    task_id = await _create_task(executor_head_client, application, inspector)

    listed = (await executor_head_client.get(f"{API}/tasks", params={"page_size": 100})).json()
    listed_ids = {row["id"] for row in listed["items"]}
    assert task_id in listed_ids

    resp = await executor_head_client.get(f"{API}/tasks/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[-1] == "ID"
    exported_ids = _ids(sheet)
    assert exported_ids == listed_ids
    assert task_id in exported_ids


async def test_tasks_export_applies_the_same_status_filter_as_the_list(
    executor_head_client, application: Application, inspector
) -> None:
    await _create_task(executor_head_client, application, inspector)  # status: assigned

    resp = await executor_head_client.get(
        f"{API}/tasks/export.xlsx", params={"status": "cancelled", "lang": "uz_latn"}
    )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_tasks_export_renders_labels_not_codes(
    executor_head_client, application: Application, inspector
) -> None:
    await _create_task(executor_head_client, application, inspector)

    resp = await executor_head_client.get(f"{API}/tasks/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == "Ruxsatnomani tekshirish"  # kind label, not "permit_inspection"
    assert row[1] == "Tayinlangan"  # status label, not "assigned"


async def test_tasks_export_truncates_at_the_cap_and_says_so(
    executor_head_client, application: Application, inspector, monkeypatch
) -> None:
    await _create_task(executor_head_client, application, inspector)
    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db, key):
        # `auth.deps.get_current_session` reads `session_idle_minutes`
        # through this same function on every request — the patch must
        # fall through to the real one for every key but ours.
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await executor_head_client.get(f"{API}/tasks/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_tasks_export_is_empty_not_403_for_a_caller_with_no_scope(
    other_inspector_client,
) -> None:
    resp = await other_inspector_client.get(f"{API}/tasks/export.xlsx")
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_tasks_export_rejects_an_unknown_language(executor_head_client) -> None:
    resp = await executor_head_client.get(f"{API}/tasks/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- Acts --------------------------------------------------------------


async def test_acts_export_holds_exactly_the_rows_the_list_shows(
    inspector_client, application: Application, default_checklist_id
) -> None:
    act_id = await _create_act(inspector_client, application, default_checklist_id)

    listed = (await inspector_client.get(f"{API}/acts", params={"page_size": 100})).json()
    listed_ids = {row["id"] for row in listed["items"]}
    assert act_id in listed_ids

    resp = await inspector_client.get(f"{API}/acts/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    assert [c.value for c in sheet[1]][-1] == "ID"
    exported_ids = _ids(sheet)
    assert exported_ids == listed_ids
    assert act_id in exported_ids


async def test_acts_export_applies_the_same_result_filter_as_the_list(
    inspector_client, application: Application, default_checklist_id
) -> None:
    await _create_act(inspector_client, application, default_checklist_id, result="compliant")

    resp = await inspector_client.get(
        f"{API}/acts/export.xlsx", params={"result": "violation", "lang": "uz_latn"}
    )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_acts_export_renders_labels_not_codes(
    inspector_client, application: Application, default_checklist_id
) -> None:
    await _create_act(inspector_client, application, default_checklist_id, result="warning")

    resp = await inspector_client.get(f"{API}/acts/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[1] == "Qoralama"  # status label ("draft"), not the raw code
    assert row[2] == "Eslatma"  # result label ("warning"), not the raw code


async def test_acts_export_truncates_at_the_cap_and_says_so(
    inspector_client, application: Application, default_checklist_id, monkeypatch
) -> None:
    await _create_act(inspector_client, application, default_checklist_id)
    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db, key):
        # `auth.deps.get_current_session` reads `session_idle_minutes`
        # through this same function on every request — the patch must
        # fall through to the real one for every key but ours.
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await inspector_client.get(f"{API}/acts/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_acts_export_is_empty_not_403_for_a_caller_with_no_scope(
    other_inspector_client,
) -> None:
    resp = await other_inspector_client.get(f"{API}/acts/export.xlsx")
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_acts_export_rejects_an_unknown_language(inspector_client) -> None:
    resp = await inspector_client.get(f"{API}/acts/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422


# --- Cases -----------------------------------------------------------------


async def test_cases_export_holds_exactly_the_rows_the_list_shows(
    db,
    executor_head_client,
    inspector_client,
    inspector,
    application: Application,
    default_checklist_id,
    vt_01,
) -> None:
    case_id = await _open_case(
        db, inspector_client, inspector, application, default_checklist_id, vt_01
    )

    listed = (await executor_head_client.get(f"{API}/cases", params={"page_size": 100})).json()
    listed_ids = {row["id"] for row in listed["items"]}
    assert case_id in listed_ids

    resp = await executor_head_client.get(f"{API}/cases/export.xlsx", params={"lang": "ru"})
    assert resp.status_code == 200
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Номер дела" and headers[-1] == "ID"
    exported_ids = _ids(sheet)
    assert exported_ids == listed_ids
    assert case_id in exported_ids


async def test_cases_export_applies_the_same_status_filter_as_the_list(
    db,
    executor_head_client,
    inspector_client,
    inspector,
    application: Application,
    default_checklist_id,
    vt_01,
) -> None:
    await _open_case(db, inspector_client, inspector, application, default_checklist_id, vt_01)

    resp = await executor_head_client.get(
        f"{API}/cases/export.xlsx", params={"status": "closed", "lang": "uz_latn"}
    )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_cases_export_renders_labels_not_codes(
    db,
    executor_head_client,
    inspector_client,
    inspector,
    application: Application,
    default_checklist_id,
    vt_01,
) -> None:
    await _open_case(db, inspector_client, inspector, application, default_checklist_id, vt_01)

    resp = await executor_head_client.get(f"{API}/cases/export.xlsx", params={"lang": "uz_latn"})
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[1] == "Ochilgan"  # status label ("opened"), not the raw code
    number = row[0]
    # the human case number comes first (service.NUMBER_PREFIX)
    assert isinstance(number, str) and number.startswith("VC-")


async def test_cases_export_truncates_at_the_cap_and_says_so(
    db,
    executor_head_client,
    inspector_client,
    inspector,
    application: Application,
    default_checklist_id,
    vt_01,
    monkeypatch,
) -> None:
    await _open_case(db, inspector_client, inspector, application, default_checklist_id, vt_01)
    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def one(db, key):
        # `auth.deps.get_current_session` reads `session_idle_minutes`
        # through this same function on every request — the patch must
        # fall through to the real one for every key but ours.
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", one)
    resp = await executor_head_client.get(f"{API}/cases/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_cases_export_is_empty_not_403_for_a_caller_with_no_scope(
    other_inspector_client,
) -> None:
    resp = await other_inspector_client.get(f"{API}/cases/export.xlsx")
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_cases_export_rejects_an_unknown_language(executor_head_client) -> None:
    resp = await executor_head_client.get(f"{API}/cases/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
