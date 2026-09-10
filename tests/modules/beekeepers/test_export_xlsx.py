"""Stage 13: `GET /beekeepers/export.xlsx` — the beekeepers register on
paper. `beekeepers` is a GLOBAL, unscoped catalog (ruling #182 — no zone at
all), so — like `report_forms` — its total can already exceed the list's
own `page_size<=100` ceiling on this shared, persistent test DB (lesson:
assert on something fresh, never an assumed-small or empty neighbourhood).
Compared on TOTAL count plus this test's own fresh row's presence, never on
full id-set equality against one capped list page."""

import io
import uuid

from openpyxl import load_workbook

from app.main import create_app
from app.modules.auth.models import UserPermission
from app.modules.beekeepers.permissions import BEEKEEPERS_MANAGE
from tests.conftest import make_client
from tests.modules.auth.test_sessions import make_session, make_user


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


def _unique_pinfl() -> str:
    # Leading digit 4 — the digit block this package's own test_router.py
    # already claims on this shared, persistent test DB.
    return f"4{uuid.uuid4().int % 10**13:013d}"


def _auth(client, token: str, csrf: str):
    client.cookies.set("session", token)
    client.cookies.set("csrf_token", csrf)
    client.headers["X-CSRF-Token"] = csrf
    return client


async def _registrar_client(db, *codes: str):
    user = await make_user(db, role_code="executor_staff")
    for code in codes:
        db.add(UserPermission(user_id=user.id, permission_code=code))
    _, token, csrf = await make_session(db, user)
    await db.commit()
    app = create_app()
    return make_client(app, lifespan=True), token, csrf


async def _create_beekeeper(client, *, certificate_no: str, full_name: str) -> str:
    resp = await client.post(
        "/api/v1/beekeepers",
        json={
            "certificate_no": certificate_no,
            "pinfl": _unique_pinfl(),
            "passport_series": "AB",
            "passport_number": "1234567",
            "full_name": full_name,
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


async def test_export_holds_exactly_the_rows_the_list_shows(db):
    ctx, token, csrf = await _registrar_client(db, BEEKEEPERS_MANAGE)
    async with ctx as client:
        _auth(client, token, csrf)
        certificate_no = f"EXP-{uuid.uuid4().hex[:8].upper()}"
        beekeeper_id = await _create_beekeeper(
            client, certificate_no=certificate_no, full_name="Export Test Beekeeper"
        )

        listed_total = (await client.get("/api/v1/beekeepers", params={"page_size": 1})).json()[
            "total"
        ]

        resp = await client.get("/api/v1/beekeepers/export.xlsx", params={"lang": "ru"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert int(resp.headers["x-export-total"]) == listed_total
        sheet = _sheet(resp.content)
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "Номер сертификата" and headers[-1] == "ID"
        exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert beekeeper_id in exported_ids


async def test_export_applies_the_same_filter_as_the_list(db):
    ctx, token, csrf = await _registrar_client(db, BEEKEEPERS_MANAGE)
    async with ctx as client:
        _auth(client, token, csrf)
        certificate_no = f"QRY-{uuid.uuid4().hex[:8].upper()}"
        beekeeper_id = await _create_beekeeper(
            client, certificate_no=certificate_no, full_name="Filter Test Beekeeper"
        )

        resp = await client.get("/api/v1/beekeepers/export.xlsx", params={"q": certificate_no})
        assert resp.status_code == 200
        exported_ids = {
            str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
        }
        assert exported_ids == {beekeeper_id}


async def test_export_renders_labels_not_codes(db):
    ctx, token, csrf = await _registrar_client(db, BEEKEEPERS_MANAGE)
    async with ctx as client:
        _auth(client, token, csrf)
        certificate_no = f"LBL-{uuid.uuid4().hex[:8].upper()}"
        beekeeper_id = await _create_beekeeper(
            client, certificate_no=certificate_no, full_name="Label Test Beekeeper"
        )

        resp = await client.get(
            "/api/v1/beekeepers/export.xlsx", params={"q": certificate_no, "lang": "uz_latn"}
        )
        row = next(
            r
            for r in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
            if r[-1] == beekeeper_id
        )
        assert row[0] == certificate_no  # the human number first
        assert row[7] == "Faol"  # the status LABEL, not "active"


async def test_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    ctx, token, csrf = await _registrar_client(db, BEEKEEPERS_MANAGE)
    async with ctx as client:
        _auth(client, token, csrf)
        for i in range(2):
            await _create_beekeeper(
                client, certificate_no=f"CAP-{uuid.uuid4().hex[:8].upper()}", full_name=f"Cap {i}"
            )

        from app.core import settings_store

        real_get_int = settings_store.get_int

        async def capped(db_, key):
            if key == "register_export_max_rows":
                return 1
            return await real_get_int(db_, key)

        monkeypatch.setattr(settings_store, "get_int", capped)

        resp = await client.get("/api/v1/beekeepers/export.xlsx")
        assert resp.status_code == 200
        total = int(resp.headers["x-export-total"])
        assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
        assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1
        assert resp.headers["x-export-rows"] == str(min(total, 1))


async def test_export_requires_the_permission(db):
    ctx, token, csrf = await _registrar_client(db)  # no grants
    async with ctx as client:
        _auth(client, token, csrf)
        resp = await client.get("/api/v1/beekeepers/export.xlsx")
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "ERR-ACL-001"


async def test_export_rejects_an_unknown_language(db):
    ctx, token, csrf = await _registrar_client(db, BEEKEEPERS_MANAGE)
    async with ctx as client:
        _auth(client, token, csrf)
        resp = await client.get("/api/v1/beekeepers/export.xlsx", params={"lang": "en"})
        assert resp.status_code == 422
