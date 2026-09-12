"""Stage 13: `GET /admin/public/appeals/export.xlsx` is the staff triage
list on paper — same `status` filter, same `public.appeals.manage` gate,
readable cells, the id last.

Fixtures go through the real transition (`POST /public/appeals` then
`POST .../status` / `.../answer`), the same pattern `test_appeals.py`
already uses, never by assigning `CitizenAppeal.status` by hand."""

import io

import pytest
from openpyxl import load_workbook

from app.core import settings_store, xlsx
from app.main import create_app
from tests.conftest import make_client
from tests.modules.auth.test_sessions import make_session, make_user
from tests.modules.public.conftest import appeals_manager, auth_client

pytestmark = pytest.mark.asyncio

API = "/api/v1"


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def _submit(client, *, subject: str, phone: str = "+998901234567") -> str:
    r = await client.post(
        f"{API}/public/appeals",
        json={
            "applicant_name": "Test Citizen",
            "contact": {"phone": phone, "email": "citizen@example.uz"},
            "subject": subject,
            "body": "Body text",
        },
    )
    assert r.status_code == 201
    return r.json()["number"]


async def test_export_holds_exactly_the_rows_the_list_shows(db):
    _, token, csrf = await appeals_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        number = await _submit(client, subject="Export test - holds the rows")

        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/admin/public/appeals", params={"status": "new"})
        assert listed.status_code == 200
        listed_ids = {item["id"] for item in listed.json()["items"]}
        assert listed_ids

        resp = await client.get(
            f"{API}/admin/public/appeals/export.xlsx", params={"status": "new", "lang": "ru"}
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert resp.headers["x-export-truncated"] == "false"
        sheet = _sheet(resp.content)
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "Номер" and headers[-1] == "ID"
        exported = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert exported == listed_ids

        submitted_numbers = {
            str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] == number
        }
        assert submitted_numbers == {number}


async def test_export_applies_the_same_filters_as_the_list(db):
    _, token, csrf = await appeals_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        number = await _submit(client, subject="Export test - to be closed")

        auth_client(client, token, csrf)
        listing = await client.get(f"{API}/admin/public/appeals", params={"status": "new"})
        appeal = next(item for item in listing.json()["items"] if item["number"] == number)
        advance = await client.post(
            f"{API}/admin/public/appeals/{appeal['id']}/status", json={"to_status": "closed"}
        )
        assert advance.status_code == 200

        resp = await client.get(f"{API}/admin/public/appeals/export.xlsx", params={"status": "new"})
        rows = _sheet(resp.content).iter_rows(min_row=2, values_only=True)
        exported = {str(row[-1]) for row in rows}
        assert appeal["id"] not in exported  # the row moved to "closed"; "new" no longer holds it

        resp_closed = await client.get(
            f"{API}/admin/public/appeals/export.xlsx", params={"status": "closed"}
        )
        rows_closed = _sheet(resp_closed.content).iter_rows(min_row=2, values_only=True)
        exported_closed = {str(row[-1]) for row in rows_closed}
        assert appeal["id"] in exported_closed


async def test_export_renders_labels_not_codes(db):
    _, token, csrf = await appeals_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        number = await _submit(client, subject="Export test - labels", phone="+998907654321")

        auth_client(client, token, csrf)
        resp = await client.get(
            f"{API}/admin/public/appeals/export.xlsx", params={"status": "new", "lang": "uz_latn"}
        )
        row = next(
            r for r in _sheet(resp.content).iter_rows(min_row=2, values_only=True) if r[0] == number
        )
        assert row[0] == number  # the human number first
        assert row[2] == "Yangi"  # the status label, not "new"
        assert row[4] == "+998907654321"  # phone, from `contact`


async def test_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await appeals_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        await _submit(client, subject="Cap test 1")
        await _submit(client, subject="Cap test 2")

        real_get_int = settings_store.get_int

        async def one(db_, key):
            if key == xlsx.CAP_SETTING:
                return 1
            return await real_get_int(db_, key)

        monkeypatch.setattr(settings_store, "get_int", one)

        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/public/appeals/export.xlsx", params={"status": "new"})
        total = int(resp.headers["x-export-total"])
        assert total > 1
        assert resp.headers["x-export-truncated"] == "true"
        assert resp.headers["x-export-rows"] == "1"
        assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_export_gets_the_same_403_the_list_gives_without_the_permission(db):
    other = await make_user(db)
    _, other_token, other_csrf = await make_session(db, other)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, other_token, other_csrf)
        listed = await client.get(f"{API}/admin/public/appeals")
        exported = await client.get(f"{API}/admin/public/appeals/export.xlsx")
    assert listed.status_code == exported.status_code == 403
    assert exported.json()["error"]["code"] == "ERR-ACL-001"


async def test_export_rejects_an_unknown_language(db):
    _, token, csrf = await appeals_manager(db)
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/admin/public/appeals/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
