"""Stage 13: `GET /notifications/export.xlsx` — the caller's own inbox on
paper. Same scope (own rows only, ruling R2), same `unread` filter, readable
cells, the id last."""

import io

from openpyxl import load_workbook

from app.main import create_app
from app.modules.notifications import service
from tests.conftest import make_client
from tests.modules.admin.test_organizations_admin import auth_client
from tests.modules.auth.test_sessions import make_session, make_user

API = "/api/v1"
EVENT = "permit.active"  # keeps an active `sms` template after 0060 (ruling #211)


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return sheet


async def _signed_in(db, **overrides):
    user = await make_user(db, **overrides)
    _, token, csrf = await make_session(db, user)
    return user, token, csrf


async def test_export_holds_exactly_the_rows_the_list_shows(db):
    mine, token, csrf = await _signed_in(db)
    other = await make_user(db)
    await service.notify(
        db, event_code=EVENT, recipient_user_id=mine.id, params={"permit_number": "P-EXP-1"}
    )
    await service.notify(
        db, event_code=EVENT, recipient_user_id=other.id, params={"permit_number": "P-EXP-2"}
    )
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(f"{API}/notifications", params={"page_size": 100})
        listed_ids = {row["id"] for row in listed.json()["items"]}

        resp = await client.get(f"{API}/notifications/export.xlsx", params={"lang": "ru"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        sheet = _sheet(resp.content)
        headers = [c.value for c in sheet[1]]
        assert headers[0] == "Код события" and headers[-1] == "ID"
        exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert exported_ids == listed_ids
        texts = [str(row[2]) for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert any("P-EXP-1" in t for t in texts)
        assert not any("P-EXP-2" in t for t in texts)  # another user's own row never leaks


async def test_export_applies_the_unread_filter(db):
    user, token, csrf = await _signed_in(db)
    rows = await service.notify(db, event_code=EVENT, recipient_user_id=user.id, params={})
    await service.notify(db, event_code=EVENT, recipient_user_id=user.id, params={})
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        await client.post(f"{API}/notifications/{rows[0].id}/read")

        resp = await client.get(f"{API}/notifications/export.xlsx", params={"unread": True})
        assert resp.status_code == 200
        exported_ids = {
            str(row[-1]) for row in _sheet(resp.content).iter_rows(min_row=2, values_only=True)
        }
        assert str(rows[0].id) not in exported_ids  # read, so filtered out


async def test_export_renders_the_status_transition_and_the_read_state(db):
    user, token, csrf = await _signed_in(db)
    await service.notify(
        db,
        event_code="application.status_changed",
        recipient_user_id=user.id,
        params={"status_from": "SUBMITTED", "status_to": "APPROVED"},
    )
    await db.commit()

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/notifications/export.xlsx", params={"lang": "uz_latn"})
        row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
        assert row[0] == "application.status_changed"  # the code, no label map
        assert row[3] == "SUBMITTED"  # status_from
        assert row[4] == "APPROVED"  # status_to
        assert row[8] is None  # read_at, never read


async def test_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    user, token, csrf = await _signed_in(db)
    for _ in range(2):
        await service.notify(db, event_code=EVENT, recipient_user_id=user.id, params={})
    await db.commit()

    from app.core import settings_store

    real_get_int = settings_store.get_int

    async def capped(db_, key):
        if key == "register_export_max_rows":
            return 1
        return await real_get_int(db_, key)

    monkeypatch.setattr(settings_store, "get_int", capped)

    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/notifications/export.xlsx")
        assert resp.status_code == 200
        total = int(resp.headers["x-export-total"])
        assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
        assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1
        assert resp.headers["x-export-rows"] == str(min(total, 1))


async def test_export_requires_authentication(db):
    async with make_client(create_app(), lifespan=True) as client:
        resp = await client.get(f"{API}/notifications/export.xlsx")
    assert resp.status_code == 401


async def test_export_rejects_an_unknown_language(db):
    _, token, csrf = await _signed_in(db)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/notifications/export.xlsx", params={"lang": "en"})
        assert resp.status_code == 422
