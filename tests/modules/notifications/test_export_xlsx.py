"""Stage 13 (ruling #204): the notification template register on paper — the
same filters, the same permission gate, readable cells, never the body."""

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.core import settings_store
from app.main import create_app
from app.modules.notifications.permissions import TEMPLATES_MANAGE
from tests.conftest import make_client
from tests.modules.admin.test_organizations_admin import auth_client, signed_in_with

pytestmark = pytest.mark.asyncio

API = "/api/v1/admin/notification-templates"


def _payload(event_code: str, channel: str = "inapp", **overrides) -> dict:
    body = {
        "event_code": event_code,
        "channel": channel,
        "body": {"uz_cyrl": "Матн {x}", "uz_latn": "Matn {x}", "ru": "Текст {x}"},
    }
    body.update(overrides)
    return body


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None  # a workbook this module wrote always has one
    return sheet


def _exported_ids(content: bytes) -> set[str]:
    return {str(row[-1]) for row in _sheet(content).iter_rows(min_row=2, values_only=True)}


async def test_export_holds_exactly_the_rows_the_list_shows(db):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    code = f"test.export{uuid.uuid4().hex[:8]}"
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(API, json=_payload(code))
        assert created.status_code == 201, created.text
        template_id = created.json()["id"]

        listed = await client.get(API, params={"event_code": code})
        resp = await client.get(f"{API}/export.xlsx", params={"event_code": code, "lang": "ru"})
    assert listed.status_code == 200, listed.text
    listed_ids = {row["id"] for row in listed.json()["items"]}
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == listed_ids == {template_id}


async def test_export_applies_the_same_filters_as_the_list(db):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    code = f"test.filter{uuid.uuid4().hex[:8]}"
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        inapp = await client.post(API, json=_payload(code, channel="inapp"))
        sms = await client.post(API, json=_payload(code, channel="sms"))
        assert inapp.status_code == 201 and sms.status_code == 201

        resp = await client.get(f"{API}/export.xlsx", params={"event_code": code, "channel": "sms"})
    assert resp.status_code == 200, resp.text
    exported_ids = _exported_ids(resp.content)
    assert exported_ids == {sms.json()["id"]}


async def test_export_renders_labels_not_codes(db):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    code = f"test.labels{uuid.uuid4().hex[:8]}"
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            API, json=_payload(code, channel="email", subject={"uz_latn": "Mavzu matni"})
        )
        assert created.status_code == 201, created.text
        resp = await client.get(
            f"{API}/export.xlsx", params={"event_code": code, "lang": "uz_latn"}
        )
    assert resp.status_code == 200, resp.text
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == code  # event code first
    assert row[1] == "Email"  # channel label
    assert row[2] == "Amaldagi"  # status label, not "active"
    assert row[4] == "Mavzu matni"  # subject, localized


async def test_export_never_exports_the_body(db):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    code = f"test.nobody{uuid.uuid4().hex[:8]}"
    secret_body = "SECRET-BODY-TEXT-SHOULD-NOT-APPEAR"
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        created = await client.post(
            API,
            json=_payload(code, body={"uz_latn": secret_body, "uz_cyrl": "x", "ru": "x"}),
        )
        assert created.status_code == 201, created.text
        resp = await client.get(f"{API}/export.xlsx", params={"event_code": code})
    assert resp.status_code == 200, resp.text
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert secret_body not in "".join(str(v) for v in row if v is not None)


async def test_export_truncates_at_the_cap_and_says_so(db, monkeypatch):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    code_prefix = f"test.cap{uuid.uuid4().hex[:8]}"
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        await client.post(API, json=_payload(f"{code_prefix}.a"))
        await client.post(API, json=_payload(f"{code_prefix}.b"))

        real_get_int = settings_store.get_int

        async def one(_db, key):
            if key == "register_export_max_rows":
                return 1
            return await real_get_int(_db, key)

        monkeypatch.setattr(settings_store, "get_int", one)
        resp = await client.get(f"{API}/export.xlsx")
    assert resp.status_code == 200, resp.text
    total = int(resp.headers["x-export-total"])
    assert total >= 2
    assert resp.headers["x-export-truncated"] == "true"
    assert resp.headers["x-export-rows"] == "1"
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) == 1


async def test_export_without_the_permission_matches_the_list_status(db):
    _, token, csrf = await signed_in_with(db)  # no grants
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        listed = await client.get(API)
        resp = await client.get(f"{API}/export.xlsx")
    assert listed.status_code == 403
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == listed.json()["error"]["code"]


async def test_export_rejects_an_unknown_language(db):
    _, token, csrf = await signed_in_with(db, TEMPLATES_MANAGE)
    await db.commit()
    async with make_client(create_app(), lifespan=True) as client:
        auth_client(client, token, csrf)
        resp = await client.get(f"{API}/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
