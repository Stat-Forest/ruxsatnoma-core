"""Stage 13 (ruling #204): the payments module's register exports are their
screen on paper — same scope, same filters, readable cells, the id last.

`/invoices/export.xlsx` (Task C.1) below; the other payments lists (Task
C.2) land in this same file, one commit each."""

import io

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.asyncio


def _sheet(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None  # a fresh Workbook always has one active sheet
    return sheet


# --- /invoices/export.xlsx (Task C.1) ---------------------------------------


async def test_invoices_export_holds_exactly_the_rows_the_list_shows(payments_view_client, invoice):
    # Scoped by `application_id` — the same filter the screen's own "search
    # by application" box sends — so the comparison is deterministic
    # regardless of what earlier tests left in this shared, persistent test
    # database (lesson: "The test DB is shared, persistent, and never empty").
    listed = (
        await payments_view_client.get(
            "/api/v1/invoices", params={"application_id": str(invoice.application_id)}
        )
    ).json()
    listed_ids = {row["id"] for row in listed["items"]}
    assert listed_ids  # non-empty: this invoice is visible to a payments.view holder

    resp = await payments_view_client.get(
        "/api/v1/invoices/export.xlsx",
        params={"application_id": str(invoice.application_id), "lang": "ru"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-export-truncated"] == "false"
    sheet = _sheet(resp.content)
    headers = [c.value for c in sheet[1]]
    assert headers[0] == "Номер счёта" and headers[-1] == "ID"
    exported_ids = {str(row[-1]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported_ids == listed_ids


async def test_invoices_export_applies_the_same_filters_as_the_list(payments_view_client, invoice):
    resp = await payments_view_client.get(
        "/api/v1/invoices/export.xlsx",
        params={"application_id": str(invoice.application_id), "status": "paid", "lang": "uz_latn"},
    )
    assert resp.status_code == 200
    assert list(_sheet(resp.content).iter_rows(min_row=2, values_only=True)) == []


async def test_invoices_export_renders_labels_not_codes(payments_view_client, invoice):
    resp = await payments_view_client.get(
        "/api/v1/invoices/export.xlsx",
        params={"application_id": str(invoice.application_id), "lang": "uz_latn"},
    )
    row = next(_sheet(resp.content).iter_rows(min_row=2, values_only=True))
    assert row[0] == invoice.number  # the human number first
    assert row[1] == "Toʻlov kutilmoqda"  # the status label, not "pending"


async def test_invoices_export_truncates_at_the_cap_and_says_so(
    payments_view_client, invoice, monkeypatch
):
    from app.core import settings_store

    original_get_int = settings_store.get_int

    async def capped(db, key):
        # `get_current_session` (every authenticated request) also reads
        # `session_idle_minutes` through this same function — only the
        # export's own cap key is overridden here.
        if key == "register_export_max_rows":
            return 1
        return await original_get_int(db, key)

    monkeypatch.setattr(settings_store, "get_int", capped)
    resp = await payments_view_client.get("/api/v1/invoices/export.xlsx")
    assert resp.status_code == 200
    total = int(resp.headers["x-export-total"])
    assert resp.headers["x-export-truncated"] == ("true" if total > 1 else "false")
    assert len(list(_sheet(resp.content).iter_rows(min_row=2, values_only=True))) <= 1


async def test_invoices_export_matches_the_list_status_for_a_caller_with_no_scope(owner_client):
    # `owner_client` with `application_id` omitted gets whatever `GET
    # /invoices` gives it today (currently 403 — `payments.view` gated,
    # backend-gaps finding 3; a later stage may change this to the owner's
    # own list, 200) — the export must answer the SAME status and, on a
    # 200, the same id set (brief shape 5).
    listed_resp = await owner_client.get("/api/v1/invoices")
    export_resp = await owner_client.get("/api/v1/invoices/export.xlsx")
    assert export_resp.status_code == listed_resp.status_code
    if listed_resp.status_code == 200:
        listed_ids = {row["id"] for row in listed_resp.json()["items"]}
        exported_ids = {
            str(row[-1])
            for row in _sheet(export_resp.content).iter_rows(min_row=2, values_only=True)
        }
        assert exported_ids == listed_ids


async def test_invoices_export_rejects_an_unknown_language(payments_view_client):
    resp = await payments_view_client.get("/api/v1/invoices/export.xlsx", params={"lang": "en"})
    assert resp.status_code == 422
