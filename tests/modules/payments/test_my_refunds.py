"""The citizen's side of refunds (stage 11, rulings R1 and R3): `GET /refunds`
lists their own, `GET /refunds/{id}` opens to the owner, and both blank the
accountant's working fields (`suggested_amount`, `suggestion_reason`,
`components`, `available_sources`) for a reader who is not staff. The staff
register and the staff single read are unchanged.

Fixtures come from `test_refunds.py` (the paid invoice with a 100-day period
whose hint is a real number — what proves the blanking blanks something)."""

import uuid

import httpx

from app.modules.applications.models import Application
from app.modules.payments.models import Invoice

# Every fixture a borrowed fixture itself depends on must be imported too, or
# pytest reports "fixture 'head' not found" from inside the borrowed one:
# `head_client` needs `head`, `refund_application` needs `published_contour`,
# and `published_contour` (applications/conftest.py) itself needs
# `contours_layer`/`approval_doc` — not already visible in this package via
# `payments/conftest.py`, which imports `leshoz` but not those two.
from tests.modules.applications.conftest import published_contour as published_contour
from tests.modules.gis.conftest import approval_doc as approval_doc
from tests.modules.gis.conftest import contours_layer as contours_layer
from tests.modules.payments.test_refunds import head as head
from tests.modules.payments.test_refunds import head_client as head_client
from tests.modules.payments.test_refunds import paid_refund_invoice as paid_refund_invoice
from tests.modules.payments.test_refunds import refund_application as refund_application
from tests.modules.payments.test_refunds import rf01 as rf01

REFUNDS = "/api/v1/refunds"
BLANKED = {
    "suggested_amount": None,
    "suggestion_reason": None,
    "components": [],
    "available_sources": [],
}


async def _file(client: httpx.AsyncClient, application_id: uuid.UUID, rf01: uuid.UUID) -> str:
    """Every caller in this file files as the owner (never staff), so the
    201 echo itself must already carry the blanked hint (stage 11 fix
    wave — `request_refund` now runs the response through `_refund_out`
    too) — asserted here once rather than in each of this file's own
    tests."""
    response = await client.post(
        REFUNDS, json={"application_id": str(application_id), "basis_item_id": str(rf01)}
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["suggested_amount"] is None
    return body["id"]


async def test_the_owner_lists_their_own_refund_with_the_hint_blanked(
    owner_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    refund_id = await _file(owner_client, refund_application.id, rf01)

    response = await owner_client.get(REFUNDS)

    assert response.status_code == 200, response.text
    rows = {item["id"]: item for item in response.json()["items"]}
    assert refund_id in rows
    assert {key: rows[refund_id][key] for key in BLANKED} == BLANKED
    assert rows[refund_id]["status"] == "requested"
    assert rows[refund_id]["due_at"]  # the 20-working-day deadline is theirs to know


async def test_the_owner_opens_their_own_refund_and_sees_no_accounting(
    owner_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    refund_id = await _file(owner_client, refund_application.id, rf01)

    response = await owner_client.get(f"{REFUNDS}/{refund_id}")

    assert response.status_code == 200, response.text
    body = response.json()
    assert {key: body[key] for key in BLANKED} == BLANKED
    assert body["invoice_id"] == str(paid_refund_invoice.id)


async def test_the_accountant_still_sees_the_hint_and_the_sources(
    owner_client: httpx.AsyncClient,
    payments_view_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    """The blanking is for the citizen, not a regression for the accountant:
    the same refund read by `payments.view` carries the real hint (the
    fixture's 100-day period prices a non-null one) and the invoice's split."""
    refund_id = await _file(owner_client, refund_application.id, rf01)

    response = await payments_view_client.get(f"{REFUNDS}/{refund_id}")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["suggested_amount"] is not None
    assert body["available_sources"] != []


async def test_a_submitted_decisions_comment_replaces_the_citizens_own(
    owner_client: httpx.AsyncClient,
    payments_view_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    """`comment` is one column three writers share:
    `backoffice_service.submit_refund_decision`/`approve_refund` both do
    `if comment: row.comment = comment`, so the accountant's own note
    overwrites whatever the citizen wrote when they filed. Once that has
    happened nobody can tell whose text `comment` holds any more, so
    `_refund_out` blanks it for the owner from the moment the refund
    leaves `requested` — the owner's own read shows `comment: null`
    (never the accountant's note), while the accountant's own read of the
    same refund still carries it."""
    refund_id = await _file(owner_client, refund_application.id, rf01)

    decision = await payments_view_client.post(
        f"{REFUNDS}/{refund_id}/submit-decision",
        json={
            "final_amount": "600000.00",
            "components": [{"recipient_id": None, "amount": "600000.00"}],
            "comment": "verified against the bank statement",
        },
    )
    assert decision.status_code == 200, decision.text

    owner_read = await owner_client.get(f"{REFUNDS}/{refund_id}")
    assert owner_read.status_code == 200, owner_read.text
    owner_body = owner_read.json()
    assert owner_body["comment"] is None
    assert owner_body["status"] == "in_review"

    staff_read = await payments_view_client.get(f"{REFUNDS}/{refund_id}")
    assert staff_read.status_code == 200, staff_read.text
    assert staff_read.json()["comment"] == "verified against the bank statement"


async def test_a_stranger_cannot_open_or_list_somebody_elses_refund(
    owner_client: httpx.AsyncClient,
    applicant_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    """404, never 403 — the same oracle rule `get_invoice_for_actor` states;
    and the stranger's own list is empty, and their `?application_id=` for
    somebody else's application is the same 404 `GET /invoices` gives."""
    refund_id = await _file(owner_client, refund_application.id, rf01)

    single = await applicant_client.get(f"{REFUNDS}/{refund_id}")
    assert single.status_code == 404, single.text
    assert single.json()["error"]["code"] == "ERR-SYS-003"

    listed = await applicant_client.get(REFUNDS)
    assert listed.status_code == 200, listed.text
    assert listed.json()["items"] == []

    by_application = await applicant_client.get(f"{REFUNDS}?application_id={refund_application.id}")
    assert by_application.status_code == 404, by_application.text


async def test_the_head_still_reads_the_register_through_payments_confirm(
    head_client: httpx.AsyncClient,
) -> None:
    """`executor_head` holds `payments.confirm` and NOT `payments.view`
    (whole-branch review Important 3 of stage 7.9) — the service gate must
    keep admitting them the way the router dependency did."""
    response = await head_client.get(REFUNDS)

    assert response.status_code == 200, response.text


async def test_the_owners_refund_export_blanks_what_their_screen_blanks(
    owner_client: httpx.AsyncClient,
    head_client: httpx.AsyncClient,
    refund_application: Application,
    paid_refund_invoice: Invoice,
    rf01: uuid.UUID,
) -> None:
    """Stage 13 on top of stage 11: the citizen's file is their list on
    paper — their own row, the accountant's hint blanked; the head's file of
    the same row carries the hint. A file that showed the citizen a figure
    their screen hides would be the leaking twin of this project's usual
    hiding defect."""
    import io

    from openpyxl import load_workbook

    refund_id = await _file(owner_client, refund_application.id, rf01)

    def sheet_rows(content: bytes) -> dict[str, dict[str, object]]:
        sheet = load_workbook(io.BytesIO(content)).active
        assert sheet is not None
        headers = [str(c.value) for c in sheet[1]]
        return {
            str(row[-1]): dict(zip(headers, (c for c in row), strict=True))
            for row in sheet.iter_rows(min_row=2, values_only=True)
        }

    mine = await owner_client.get(f"{REFUNDS}/export.xlsx", params={"lang": "ru"})
    assert mine.status_code == 200, mine.text
    listed = {item["id"] for item in (await owner_client.get(REFUNDS)).json()["items"]}
    rows = sheet_rows(mine.content)
    assert set(rows) == listed and refund_id in rows
    assert rows[refund_id]["Рекомендовано"] is None

    theirs = await head_client.get(
        f"{REFUNDS}/export.xlsx",
        params={"lang": "ru", "application_id": str(refund_application.id)},
    )
    assert theirs.status_code == 200, theirs.text
    staff_rows = sheet_rows(theirs.content)
    assert refund_id in staff_rows
    assert staff_rows[refund_id]["Рекомендовано"] is not None  # the hint, computed at filing
