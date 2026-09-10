"""Generic register → .xlsx renderer for the export routes (stage 13, ruling #204).

Level 0: knows columns, cells and languages — nothing about applications or
invoices. A module describes its register as `Column`s in its own `export.py`
and calls `render`; its router wraps the bytes with `xlsx_response`.

Every `datetime` is written in Asia/Tashkent (`backend/CLAUDE.md` "Time":
store UTC, display Tashkent — a spreadsheet IS display) and NAIVE, because
Excel has no timezone-aware cell. `Decimal` and `date` go through as they
are — openpyxl writes a numeric and a date cell. A missing value is an empty
cell, never the string "None".

The last column is always the row's UUID (`id_column`, ruling R4): two rows
that read alike in every visible column can still be told apart, and the
renderer refuses a column set that forgets it rather than trusting each of
the twenty-odd modules to remember.
"""

import io
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Literal

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core import files
from app.core.time import TASHKENT

Lang = Literal["uz_latn", "ru"]
LANGS: tuple[Lang, ...] = ("uz_latn", "ru")
CellValue = str | int | float | Decimal | date | datetime | bool | None
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# The `system_settings` key every export route reads its row cap from — a
# key of its own, not the prosecutor's `search_export_max_rows`: his figure
# is a legal one (ruling #20), this one is operational (ruling R3).
CAP_SETTING = "register_export_max_rows"
_EXPOSED_HEADERS = "X-Export-Total, X-Export-Rows, X-Export-Truncated"
_SHEET_TITLE_MAX = 31  # Excel's own cap on a sheet name


@dataclass(frozen=True)
class Column[R]:
    """One export column: a stable ASCII `key`, a header per language
    (both of `LANGS`, checked at render time), and the accessor that reads
    the cell off a row."""

    key: str
    header: dict[str, str]
    value: Callable[[R], CellValue]
    width: int = 16


def id_column(header: dict[str, str] | None = None) -> Column[Any]:
    """The mandatory LAST column of every register export (ruling R4).
    `Column[Any]`: every row type this reads has an `id`, and `Column` is
    contravariant in its row (the row only ever enters `value`)."""
    return Column(
        key="id",
        header=header or {"uz_latn": "ID", "ru": "ID"},
        value=lambda r: str(r.id),
        width=38,
    )


def _cell(value: CellValue) -> CellValue:
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(TASHKENT)
        return value.replace(tzinfo=None)
    return value


def render[R](rows: Sequence[R], columns: Sequence[Column[R]], *, lang: Lang, title: str) -> bytes:
    """One sheet: a bold header row in `lang`, one row per item, the header
    frozen so it stays visible while scrolling."""
    for column in columns:
        missing = [code for code in LANGS if code not in column.header]
        if missing:
            raise ValueError(f"column {column.key!r} has no header for {missing}")
    if not columns or columns[-1].key != "id":
        raise ValueError("the last export column must be the id column (ruling R4)")

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # a fresh Workbook always has one active sheet
    sheet.title = title[:_SHEET_TITLE_MAX]
    sheet.append([column.header[lang] for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_cell(column.value(row)) for column in columns])
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xlsx_response(data: bytes, *, filename: str, total: int, cap: int) -> Response:
    """The file plus the three headers the adminka reads to warn about a cut
    (ruling R3). `Access-Control-Expose-Headers`, because the adminka is on
    another origin and a browser hides every non-simple response header
    from it otherwise — a truncation the client cannot see would be the
    hiding kind of defect this project keeps finding."""
    return Response(
        content=data,
        media_type=MEDIA_TYPE,
        headers={
            "Content-Disposition": files.content_disposition("attachment", filename),
            "X-Content-Type-Options": "nosniff",
            "X-Export-Total": str(total),
            "X-Export-Rows": str(min(total, cap)),
            "X-Export-Truncated": "true" if total > cap else "false",
            "Access-Control-Expose-Headers": _EXPOSED_HEADERS,
        },
    )
