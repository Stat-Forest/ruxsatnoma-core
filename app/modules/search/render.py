"""PDF and XLSX rendering for a search export (С22, decision #98).

Reuses the tool `permits/render.py` and `reports/render.py` both already use
for PDF (WeasyPrint) rather than adding a second mechanism — but not
`permits`'s own PDF/A hardening (embedded fonts, font-substitution guard,
deterministic bytes for a signature to anchor to): that machinery exists for
the document four ERI signatures are taken over, this one is a plain
administrative export nobody signs, exactly the shape `reports/render.py`
already established for its own table-to-PDF export.

The watermark is rendered INTO the document, never attached as a header
field (С22 track brief, verbatim): the PDF gets a `position: fixed` element
that WeasyPrint repeats on every page of paginated output, and the XLSX gets
a real, merged first row every sheet carries. Both are content a viewer sees
on opening the file, not a filename or a document-properties field nobody
reads.
"""

import html
import io
import os
import sys
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Row

from app.core.time import TASHKENT

# The exact five-line fix `permits/render.py` and `reports/render.py` each
# carry, duplicated rather than imported for the same reason theirs is:
# cross-module, non-service imports are not allowed here, so a third
# independently-imported module doing `import weasyprint` needs its own copy
# of the guard. setdefault, so an operator's own value wins; no effect on
# Linux, in CI or in the container.
if sys.platform == "darwin":  # pragma: no cover - a developer-machine path
    os.environ.setdefault("DYLD_FALLBACK_LIBRARY_PATH", "/opt/homebrew/lib:/usr/local/lib:/usr/lib")

from weasyprint import HTML  # noqa: E402 - must follow the dyld fix-up above

_KIND_TITLES = {"applications": "Заявки", "permits": "Разрешения"}
_HEADERS = ("№", "Заявитель", "Организация", "Статус", "Дата создания")


def _org_name(name: dict[str, Any] | None) -> str:
    """Same fallback order `reports/render.py::_label` uses: `ru` first (this
    export's own administrative audience), then whichever Uzbek script the
    row actually has — `uz_latn` is the one decision #90 guarantees on every
    row, `uz_cyrl` is not."""
    if not name:
        return ""
    return str(name.get("ru") or name.get("uz_cyrl") or name.get("uz_latn") or "")


def _row_cells(row: Row) -> tuple[str, str, str, str, str]:
    """One export row as five display strings, in `_HEADERS`'s own order.
    `created_at` is stored UTC and DISPLAYED Asia/Tashkent (`backend/
    CLAUDE.md` "Time" rule) — this is a value printed on a document a human
    reads, exactly the case that rule is for."""
    return (
        str(row.number or ""),
        str(row.applicant_name or ""),
        _org_name(row.organization_name),
        str(row.status),
        row.created_at.astimezone(TASHKENT).strftime("%Y-%m-%d %H:%M"),
    )


def watermark_text(full_name: str, exported_on: date) -> str:
    """ФИО + дата, exactly what ruling #20 names — one string both renderers
    stamp into the document, so a test asserting "the watermark carries the
    operator's name and the date" checks this single formatting rather than
    two independent ones that could drift apart."""
    return f"{full_name} — {exported_on.isoformat()}"


def render_pdf(rows: list[Row], *, kind: str, watermark: str) -> bytes:
    """A plain HTML table -> PDF, default fonts, no PDF/A — landscape so the
    organization column has room. The watermark `<div>` is `position: fixed`
    outside the table flow, which WeasyPrint repeats on every generated page
    (the same technique used to fake a running header/footer where CSS paged
    media's own `@page` margin boxes cannot hold arbitrary HTML)."""
    header_cells = "".join(f"<th>{html.escape(h)}</th>" for h in _HEADERS)
    body_rows: list[str] = []
    for row in rows:
        cells = "".join(f"<td>{html.escape(c)}</td>" for c in _row_cells(row))
        body_rows.append(f"<tr>{cells}</tr>")

    title = html.escape(_KIND_TITLES.get(kind, kind))
    mark = html.escape(watermark)
    document = f"""<!doctype html>
<html><head><meta charset="utf-8"><style>
@page {{ size: A4 landscape; margin: 15mm; }}
body {{ font-family: sans-serif; }}
h1 {{ font-size: 14px; margin: 0 0 8px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9px; }}
th, td {{ border: 1px solid #333; padding: 3px 5px; text-align: left; }}
th {{ background: #eee; }}
.watermark {{
  position: fixed;
  top: 42%;
  left: 10%;
  font-size: 40px;
  color: rgba(120, 0, 0, 0.18);
  transform: rotate(-30deg);
  white-space: nowrap;
  z-index: 1000;
}}
</style></head>
<body>
<div class="watermark">{mark}</div>
<h1>{title}</h1>
<table><thead><tr>{header_cells}</tr></thead><tbody>{"".join(body_rows)}</tbody></table>
</body></html>"""

    pdf = HTML(string=document).write_pdf()
    assert pdf is not None
    return pdf


def render_xlsx(rows: list[Row], *, kind: str, watermark: str) -> bytes:
    """One sheet: a merged watermark row, a header row, one row per result."""
    workbook = Workbook()
    sheet = workbook.active
    # A freshly-constructed Workbook() always has one active sheet (same
    # narrowing `reports/render.py::render_excel` documents for its own
    # `.active` call).
    assert sheet is not None
    sheet.title = _KIND_TITLES.get(kind, kind)[:31]  # Excel's own sheet-name cap

    sheet.append([watermark])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    _watermark_cell(sheet).font = Font(italic=True, color="999999")
    _watermark_cell(sheet).alignment = Alignment(horizontal="center")

    sheet.append(list(_HEADERS))
    for row in rows:
        sheet.append(list(_row_cells(row)))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _watermark_cell(sheet: Worksheet) -> Any:
    return sheet.cell(row=1, column=1)
